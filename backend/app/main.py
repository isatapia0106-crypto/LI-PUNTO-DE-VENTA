import os
import json
import secrets
from datetime import date, datetime, timezone, timedelta
from zoneinfo import ZoneInfo
from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import jwt
from jwt.exceptions import InvalidTokenError
from pwdlib import PasswordHash
from fastapi import FastAPI, HTTPException, Header, Depends, Request
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, ConfigDict, field_validator
from sqlalchemy import create_engine, ForeignKey, String, Integer, Numeric, DateTime, select, UniqueConstraint, inspect, or_, and_, Index, text, event
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, relationship, Session

DATABASE_URL = os.getenv('DATABASE_URL', 'sqlite:///./pos.db')
engine = create_engine(DATABASE_URL, connect_args={'check_same_thread': False} if DATABASE_URL.startswith('sqlite') else {}, pool_pre_ping=True)

if DATABASE_URL.startswith('sqlite'):
    @event.listens_for(engine, 'connect')
    def sqlite_connect(connection, record):
        connection.isolation_level = None
        connection.execute('PRAGMA foreign_keys=ON')
        connection.execute('PRAGMA busy_timeout=15000')

    @event.listens_for(engine, 'begin')
    def sqlite_begin(connection):
        connection.exec_driver_sql('BEGIN IMMEDIATE')

class Base(DeclarativeBase):
    pass

class Branch(Base):
    __tablename__ = 'branches'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    name: Mapped[str] = mapped_column(String(100))
    __table_args__ = (UniqueConstraint('empresa_id', 'name'),)

class User(Base):
    __tablename__ = 'users'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    username: Mapped[str] = mapped_column(String(80), unique=True)
    password_hash: Mapped[str] = mapped_column(String(300))
    role: Mapped[str] = mapped_column(String(40))
    active: Mapped[bool] = mapped_column(default=True)
    token_version: Mapped[int] = mapped_column(Integer, default=0, server_default='0')

class UserBranch(Base):
    __tablename__ = 'user_branches'
    user_id: Mapped[int] = mapped_column(ForeignKey('users.id'), primary_key=True)
    branch_id: Mapped[int] = mapped_column(ForeignKey('branches.id'), primary_key=True)

class CashRegister(Base):
    __tablename__ = 'cash_registers'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    branch_id: Mapped[int] = mapped_column(ForeignKey('branches.id'))
    name: Mapped[str] = mapped_column(String(80))
    __table_args__ = (UniqueConstraint('branch_id', 'name'),)

class CashSession(Base):
    __tablename__ = 'cash_sessions'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer)
    branch_id: Mapped[int] = mapped_column(ForeignKey('branches.id'))
    register_id: Mapped[int | None] = mapped_column(ForeignKey('cash_registers.id'), nullable=True)
    cashier_id: Mapped[int | None] = mapped_column(ForeignKey('users.id'), nullable=True)
    close_snapshot: Mapped[str | None] = mapped_column(String(20000), nullable=True)
    expected_on_close: Mapped[Decimal | None] = mapped_column(Numeric(12, 2), nullable=True)
    closed_by: Mapped[int | None] = mapped_column(ForeignKey('users.id'), nullable=True)
    opening: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    counted: Mapped[Decimal | None] = mapped_column(Numeric(12, 2), nullable=True)
    status: Mapped[str] = mapped_column(String(10), default='open')
    opened_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    closed_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)

Index('uq_open_register', CashSession.register_id, unique=True, sqlite_where=text("status = 'open'"), postgresql_where=text("status = 'open'"))
Index('uq_open_cashier_branch', CashSession.branch_id, CashSession.cashier_id, unique=True, sqlite_where=text("status = 'open'"), postgresql_where=text("status = 'open'"))

class CashMovement(Base):
    __tablename__ = 'cash_movements'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer)
    branch_id: Mapped[int] = mapped_column(Integer)
    session_id: Mapped[int] = mapped_column(ForeignKey('cash_sessions.id'))
    amount: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    actor_id: Mapped[int | None] = mapped_column(ForeignKey('users.id'), nullable=True)
    request_key: Mapped[str | None] = mapped_column(String(100), nullable=True)
    __table_args__ = (UniqueConstraint('empresa_id', 'request_key'),)
    kind: Mapped[str] = mapped_column(String(20))
    reason: Mapped[str] = mapped_column(String(160))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class StockMovement(Base):
    __tablename__ = 'stock_movements'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer)
    branch_id: Mapped[int] = mapped_column(Integer)
    product_id: Mapped[int] = mapped_column(ForeignKey('products.id'))
    actor_id: Mapped[int | None] = mapped_column(ForeignKey('users.id'), nullable=True)
    change: Mapped[int] = mapped_column(Integer)
    reason: Mapped[str] = mapped_column(String(30))
    reference_id: Mapped[int | None] = mapped_column(Integer, nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class StockTransfer(Base):
    __tablename__ = 'stock_transfers'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    source_branch_id: Mapped[int] = mapped_column(ForeignKey('branches.id'))
    target_branch_id: Mapped[int] = mapped_column(ForeignKey('branches.id'))
    product_id: Mapped[int] = mapped_column(ForeignKey('products.id'))
    quantity: Mapped[int] = mapped_column(Integer)
    request_key: Mapped[str] = mapped_column(String(100))
    actor_id: Mapped[int] = mapped_column(ForeignKey('users.id'))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    __table_args__ = (UniqueConstraint('empresa_id', 'request_key'),)

class Product(Base):
    __tablename__ = 'products'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    sku: Mapped[str] = mapped_column(String(60))
    name: Mapped[str] = mapped_column(String(160))
    price: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    barcode: Mapped[str | None] = mapped_column(String(80), nullable=True)
    unit: Mapped[str] = mapped_column(String(30), default='pieza')
    tax_rate: Mapped[Decimal] = mapped_column(Numeric(6, 4), default=Decimal('.16'))
    tax_exempt: Mapped[bool] = mapped_column(default=False)
    price_includes_tax: Mapped[bool] = mapped_column(default=False)
    active: Mapped[bool] = mapped_column(default=True)
    __table_args__ = (UniqueConstraint('empresa_id', 'sku'), UniqueConstraint('empresa_id', 'barcode'))

class Customer(Base):
    __tablename__ = 'customers'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    name: Mapped[str] = mapped_column(String(160))
    phone: Mapped[str | None] = mapped_column(String(30), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class Stock(Base):
    __tablename__ = 'stock'
    id: Mapped[int] = mapped_column(primary_key=True)
    product_id: Mapped[int] = mapped_column(ForeignKey('products.id'))
    branch_id: Mapped[int] = mapped_column(Integer)
    quantity: Mapped[int] = mapped_column(Integer, default=0)
    minimum: Mapped[int] = mapped_column(Integer, default=0, server_default='0')
    average_cost: Mapped[Decimal] = mapped_column(Numeric(12, 4), default=Decimal('0'))
    __table_args__ = (UniqueConstraint('product_id', 'branch_id'),)

class PaymentIntent(Base):
    __tablename__ = 'payment_intents'
    id: Mapped[str] = mapped_column(String(60), primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    branch_id: Mapped[int] = mapped_column(ForeignKey('branches.id'))
    actor_id: Mapped[int] = mapped_column(ForeignKey('users.id'))
    payload: Mapped[str] = mapped_column(String(20000))
    amount: Mapped[Decimal] = mapped_column(Numeric(12,2))
    status: Mapped[str] = mapped_column(String(40), default='creating')
    delivery_cash_session_id: Mapped[int | None] = mapped_column(Integer, nullable=True)
    delivered_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    cancel_requested_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    cancelled_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    cancel_actor_id: Mapped[int | None] = mapped_column(Integer, nullable=True)
    cancel_reason: Mapped[str | None] = mapped_column(String(300), nullable=True)
    reservation_active: Mapped[bool] = mapped_column(default=False, server_default='0')
    price_snapshot: Mapped[str | None] = mapped_column(String(20000), nullable=True)
    resolved_by: Mapped[int | None] = mapped_column(Integer, nullable=True)
    resolved_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    review_reason: Mapped[str | None] = mapped_column(String(500), nullable=True)
    preference_id: Mapped[str | None] = mapped_column(String(100), nullable=True)
    checkout_url: Mapped[str | None] = mapped_column(String(2000), nullable=True)
    payment_id: Mapped[str | None] = mapped_column(String(100), nullable=True, unique=True)
    request_key: Mapped[str] = mapped_column(String(100))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    __table_args__ = (UniqueConstraint('empresa_id', 'request_key'),)

class PaymentObservation(Base):
    __tablename__ = 'payment_observations'
    payment_id: Mapped[str] = mapped_column(String(100), primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    intent_id: Mapped[str] = mapped_column(ForeignKey('payment_intents.id'))
    status: Mapped[str] = mapped_column(String(40))
    amount: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    refunded: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


def record_payment_observation(db, intent, payment):
    pid=str(payment['id'])
    record=db.get(PaymentObservation,pid)
    if record and (record.empresa_id!=intent.empresa_id or record.intent_id!=intent.id):
        raise HTTPException(409,'Pago ya observado en otra operación')
    if record is None:
        record=PaymentObservation(payment_id=pid,empresa_id=intent.empresa_id,intent_id=intent.id)
        db.add(record)
    record.status=str(payment.get('status','unknown'))[:40]
    record.amount=Decimal(str(payment['transaction_amount']))
    record.refunded=Decimal(str(payment.get('transaction_amount_refunded',0)))
    record.updated_at=datetime.now(timezone.utc)


class PaymentRefund(Base):
    __tablename__ = 'payment_refunds'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    intent_id: Mapped[str] = mapped_column(ForeignKey('payment_intents.id'))
    payment_id: Mapped[str] = mapped_column(String(100), unique=True)
    amount: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    provider_key: Mapped[str] = mapped_column(String(60), unique=True)
    provider_refund_id: Mapped[str | None] = mapped_column(String(100), nullable=True)
    status: Mapped[str] = mapped_column(String(30), default='prepared')
    reason: Mapped[str] = mapped_column(String(300))
    actor_id: Mapped[int] = mapped_column(ForeignKey('users.id'))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    confirmed_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)


def fully_refunded(payment):
    total = Decimal(str(payment.get('transaction_amount', 0)))
    return (payment.get('status') in ('approved','refunded') and total > 0
        and Decimal(str(payment.get('transaction_amount_refunded', 0))) == total)


def reserved_quantity(db, branch_id, product_id, exclude=None):
    """Called under the branch lock for every inventory reduction."""
    intents = db.scalars(select(PaymentIntent).where(PaymentIntent.branch_id == branch_id,
        PaymentIntent.reservation_active == True))
    return sum(line['quantity'] for intent in intents if intent.id != exclude
        for line in json.loads(intent.payload)['items'] if line['product_id'] == product_id)


class Sale(Base):
    __tablename__ = 'sales'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    branch_id: Mapped[int] = mapped_column(Integer)
    customer_id: Mapped[int | None] = mapped_column(ForeignKey('customers.id'), nullable=True)
    customer_name: Mapped[str | None] = mapped_column(String(160), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    folio: Mapped[str | None] = mapped_column(String(60), nullable=True)
    discount_percent: Mapped[Decimal] = mapped_column(Numeric(5, 2), default=Decimal('0'))
    discount_total: Mapped[Decimal] = mapped_column(Numeric(12, 2), default=Decimal('0'))
    discount_reason: Mapped[str | None] = mapped_column(String(160), nullable=True)
    discount_approved_by: Mapped[int | None] = mapped_column(ForeignKey('users.id'), nullable=True)
    actor_id: Mapped[int | None] = mapped_column(ForeignKey('users.id'), nullable=True)
    subtotal: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    tax: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    total: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    payment_method: Mapped[str] = mapped_column(String(30))
    cash_session_id: Mapped[int] = mapped_column(ForeignKey('cash_sessions.id'))
    request_key: Mapped[str] = mapped_column(String(100))
    paid: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    items: Mapped[list['SaleItem']] = relationship(cascade='all, delete-orphan')
    __table_args__ = (UniqueConstraint('empresa_id', 'request_key'), UniqueConstraint('empresa_id', 'folio'))

class SaleItem(Base):
    __tablename__ = 'sale_items'
    id: Mapped[int] = mapped_column(primary_key=True)
    sale_id: Mapped[int] = mapped_column(ForeignKey('sales.id'))
    product_id: Mapped[int] = mapped_column(ForeignKey('products.id'))
    name: Mapped[str] = mapped_column(String(160))
    quantity: Mapped[int] = mapped_column(Integer)
    unit_price: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    unit: Mapped[str] = mapped_column(String(30), default='pieza')
    tax_rate: Mapped[Decimal] = mapped_column(Numeric(6, 4), default=Decimal('.16'))
    tax_exempt: Mapped[bool] = mapped_column(default=False)
    price_includes_tax: Mapped[bool] = mapped_column(default=False)
    discount_amount: Mapped[Decimal] = mapped_column(Numeric(12, 2), default=Decimal('0'))
    net_amount: Mapped[Decimal | None] = mapped_column(Numeric(12, 2), nullable=True)
    tax_amount: Mapped[Decimal | None] = mapped_column(Numeric(12, 2), nullable=True)
    cost: Mapped[Decimal] = mapped_column(Numeric(12, 4), default=Decimal('0'))

class SaleReturn(Base):
    __tablename__ = 'sale_returns'
    status: Mapped[str] = mapped_column(String(30), default='completed', server_default='completed')
    requested_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    finalized_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    provider_payment_id: Mapped[str | None] = mapped_column(String(100), nullable=True)
    provider_key: Mapped[str | None] = mapped_column(String(60), nullable=True, unique=True)
    provider_refund_id: Mapped[str | None] = mapped_column(String(100), nullable=True)
    provider_refunded_before: Mapped[Decimal | None] = mapped_column(Numeric(12,2), nullable=True)
    kind: Mapped[str] = mapped_column(String(20), default='return', server_default='return')
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    sale_id: Mapped[int] = mapped_column(ForeignKey('sales.id'))
    actor_id: Mapped[int] = mapped_column(ForeignKey('users.id'))
    cash_session_id: Mapped[int | None] = mapped_column(ForeignKey('cash_sessions.id'), nullable=True)
    reason: Mapped[str] = mapped_column(String(160))
    payment_reference: Mapped[str | None] = mapped_column(String(100), nullable=True)
    payload: Mapped[str] = mapped_column(String(20000))
    total: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    request_key: Mapped[str] = mapped_column(String(100))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    items: Mapped[list['SaleReturnItem']] = relationship(cascade='all, delete-orphan')
    __table_args__ = (UniqueConstraint('empresa_id', 'request_key'),)

class SaleReturnItem(Base):
    __tablename__ = 'sale_return_items'
    id: Mapped[int] = mapped_column(primary_key=True)
    return_id: Mapped[int] = mapped_column(ForeignKey('sale_returns.id'))
    sale_item_id: Mapped[int] = mapped_column(ForeignKey('sale_items.id'))
    quantity: Mapped[int] = mapped_column(Integer)
    restock: Mapped[bool] = mapped_column(default=True)
    total: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    __table_args__ = (UniqueConstraint('return_id', 'sale_item_id'),)

class Supplier(Base):
    __tablename__ = 'suppliers'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    name: Mapped[str] = mapped_column(String(160))
    phone: Mapped[str | None] = mapped_column(String(30), nullable=True)
    reference: Mapped[str | None] = mapped_column(String(100), nullable=True)

class Purchase(Base):
    __tablename__ = 'purchases'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer, index=True)
    branch_id: Mapped[int] = mapped_column(ForeignKey('branches.id'))
    supplier_id: Mapped[int] = mapped_column(ForeignKey('suppliers.id'))
    reference: Mapped[str] = mapped_column(String(100))
    actor_id: Mapped[int] = mapped_column(ForeignKey('users.id'))
    request_key: Mapped[str] = mapped_column(String(100))
    status: Mapped[str] = mapped_column(String(20), default='ordered')
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    items: Mapped[list['PurchaseItem']] = relationship(cascade='all, delete-orphan')
    __table_args__ = (UniqueConstraint('empresa_id', 'request_key'),)

class PurchaseItem(Base):
    __tablename__ = 'purchase_items'
    id: Mapped[int] = mapped_column(primary_key=True)
    purchase_id: Mapped[int] = mapped_column(ForeignKey('purchases.id'))
    product_id: Mapped[int] = mapped_column(ForeignKey('products.id'))
    name: Mapped[str] = mapped_column(String(160))
    quantity: Mapped[int] = mapped_column(Integer)
    received: Mapped[int] = mapped_column(Integer, default=0)
    unit_cost: Mapped[Decimal] = mapped_column(Numeric(12, 4))
    __table_args__ = (UniqueConstraint('purchase_id', 'product_id'),)

class PurchaseReceipt(Base):
    __tablename__ = 'purchase_receipts'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer)
    purchase_id: Mapped[int] = mapped_column(ForeignKey('purchases.id'))
    actor_id: Mapped[int] = mapped_column(ForeignKey('users.id'))
    request_key: Mapped[str] = mapped_column(String(100))
    payload: Mapped[str] = mapped_column(String(4000))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    __table_args__ = (UniqueConstraint('empresa_id', 'request_key'),)

class InventoryCount(Base):
    __tablename__ = 'inventory_counts'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer)
    branch_id: Mapped[int] = mapped_column(ForeignKey('branches.id'))
    product_id: Mapped[int] = mapped_column(ForeignKey('products.id'))
    actor_id: Mapped[int] = mapped_column(ForeignKey('users.id'))
    expected: Mapped[int] = mapped_column(Integer)
    counted: Mapped[int] = mapped_column(Integer)
    reason: Mapped[str] = mapped_column(String(160))
    request_key: Mapped[str] = mapped_column(String(100))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    __table_args__ = (UniqueConstraint('empresa_id', 'request_key'),)

class Audit(Base):
    __tablename__ = 'audit_logs'
    id: Mapped[int] = mapped_column(primary_key=True)
    empresa_id: Mapped[int] = mapped_column(Integer)
    branch_id: Mapped[int] = mapped_column(Integer)
    action: Mapped[str] = mapped_column(String(50))
    record_id: Mapped[int] = mapped_column(Integer)
    actor_id: Mapped[int | None] = mapped_column(Integer, nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class Line(BaseModel):
    product_id: int
    quantity: int = Field(gt=0, le=10000)

class ApprovalIn(BaseModel):
    username: str
    password: str

class SaleIn(BaseModel):
    payment_intent_id: str | None = Field(default=None, max_length=60)
    branch_id: int = Field(gt=0)
    customer_id: int | None = Field(default=None, gt=0)
    cash_session_id: int | None = Field(default=None, gt=0)
    discount_percent: Decimal = Field(default=Decimal('0'), ge=0, lt=100, decimal_places=2)
    discount_reason: str | None = Field(default=None, max_length=160)
    approval: ApprovalIn | None = None
    items: list[Line] = Field(min_length=1)
    payment_method: str
    paid: Decimal = Field(ge=0, le=9999999999, decimal_places=2)

class ProductFields(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    sku: str = Field(min_length=1, max_length=60)
    name: str = Field(min_length=1, max_length=160)
    price: Decimal = Field(gt=0, le=9999999999, decimal_places=2)
    barcode: str | None = Field(default=None, max_length=80)
    unit: str = Field(default='pieza', min_length=1, max_length=30)
    tax_rate: Decimal = Field(default=Decimal('.16'), ge=0, le=1, decimal_places=4)
    tax_exempt: bool = False
    price_includes_tax: bool = False
    active: bool = True
    @field_validator('barcode')
    @classmethod
    def empty_barcode(cls, value):
        return value or None

class ProductUpdate(ProductFields):
    branch_id: int = Field(gt=0)

class ProductIn(ProductFields):
    initial_cost: Decimal = Field(default=Decimal('0'), ge=0, le=99999999, decimal_places=4)
    branch_id: int = Field(gt=0)
    stock: int = Field(ge=0)

class CustomerIn(BaseModel):
    branch_id: int = Field(gt=0)
    name: str = Field(min_length=2, max_length=160)
    phone: str | None = Field(default=None, max_length=30)

class CustomerUpdate(CustomerIn):
    pass

class RegisterIn(BaseModel):
    branch_id: int = Field(gt=0)
    name: str = Field(min_length=1, max_length=80)

class OpenCash(BaseModel):
    register_id: int | None = Field(default=None, gt=0)
    branch_id: int = Field(gt=0)
    opening: Decimal = Field(ge=0, le=9999999999, decimal_places=2)

class CloseCash(BaseModel):
    counted: Decimal = Field(ge=0, le=9999999999, decimal_places=2)

class WithdrawCash(BaseModel):
    amount: Decimal = Field(gt=0, le=9999999999, decimal_places=2)
    reason: str = Field(min_length=3, max_length=160)

class AdjustStock(BaseModel):
    branch_id: int = Field(gt=0)
    change: int = Field(ge=-100000, le=100000)
    reason: str = Field(min_length=3, max_length=100)

class TransferStock(BaseModel):
    source_branch_id: int = Field(gt=0)
    target_branch_id: int = Field(gt=0)
    product_id: int = Field(gt=0)
    quantity: int = Field(gt=0, le=100000)

class LoginIn(BaseModel):
    username: str
    password: str

class UserIn(BaseModel):
    username: str = Field(min_length=3, max_length=80)
    password: str = Field(min_length=12, max_length=200)
    role: str
    branch_ids: list[int] = Field(min_length=1)

app = FastAPI(title='LI Punto de Venta · MVP')
password_hash = PasswordHash.recommended()
bearer = HTTPBearer(auto_error=False)
JWT_SECRET = os.getenv('JWT_SECRET')
if not JWT_SECRET:
    if os.getenv('APP_ENV') == 'production':
        raise RuntimeError('JWT_SECRET es obligatorio en producción')
    JWT_SECRET = secrets.token_urlsafe(48)  # Transient local key; restarts invalidate sessions.

ROLE_ACTIONS = {
    'admin_general': {'sale_cancel', 'sale_return', 'discount', 'purchase_write', 'purchase_read', 'count_write', 'cash_deposit', 'sale', 'catalog_write', 'customer_read', 'customer_write', 'stock_write', 'cash_open', 'cash_close', 'cash_withdraw', 'report', 'users_write'},
    'admin_sucursal': {'sale_cancel', 'sale_return', 'discount', 'purchase_write', 'purchase_read', 'count_write', 'cash_deposit', 'sale', 'catalog_write', 'customer_read', 'customer_write', 'stock_write', 'cash_open', 'cash_close', 'cash_withdraw', 'report'},
    'cajero': {'sale', 'customer_read', 'customer_write', 'cash_open', 'cash_close'},
    'almacenista': {'stock_write', 'purchase_read', 'purchase_receive'},
    'supervisor_inventarios': {'stock_write', 'report', 'purchase_read', 'purchase_receive', 'count_write'},
    'contabilidad': {'report', 'customer_read', 'purchase_read'},
    'repartidor': set(),
    'auditoria': {'report'},
}

def require(user: User, action: str):
    if action not in ROLE_ACTIONS.get(user.role, set()):
        raise HTTPException(403, 'Permiso insuficiente')

def require_any(user: User, *actions: str):
    if not any(action in ROLE_ACTIONS.get(user.role, set()) for action in actions):
        raise HTTPException(403, 'Permiso insuficiente')

def identity(credentials: HTTPAuthorizationCredentials | None = Depends(bearer)) -> User:
    if credentials is None:
        raise HTTPException(401, 'Se requiere iniciar sesión', headers={'WWW-Authenticate': 'Bearer'})
    try:
        payload = jwt.decode(credentials.credentials, JWT_SECRET, algorithms=['HS256'], options={'require': ['sub', 'exp']})
        user_id = int(payload['sub'])
    except (InvalidTokenError, ValueError, TypeError, KeyError):
        raise HTTPException(401, 'Sesión inválida', headers={'WWW-Authenticate': 'Bearer'})
    with Session(engine) as db:
        user = db.get(User, user_id)
        if not user or not user.active or payload.get('ver', 0) != user.token_version:
            raise HTTPException(401, 'Sesión inválida')
        return user

def token_for(user: User) -> str:
    return jwt.encode({'sub': str(user.id), 'ver': user.token_version, 'exp': datetime.now(timezone.utc) + timedelta(hours=8)}, JWT_SECRET, algorithm='HS256')

@app.post('/api/auth/login')
def login(data: LoginIn):
    with Session(engine) as db:
        user = db.scalar(select(User).where(User.username == data.username))
        if not user or not user.active or not password_hash.verify(data.password, user.password_hash):
            raise HTTPException(401, 'Credenciales inválidas')
        return {'access_token': token_for(user), 'token_type': 'bearer', 'role': user.role}

@app.get('/api/auth/me')
def me(user: User = Depends(identity)):
    return {'id': user.id, 'username': user.username, 'role': user.role, 'permissions': sorted(ROLE_ACTIONS.get(user.role, set()))}

@app.post('/api/users', status_code=201)
def create_user(data: UserIn, user: User = Depends(identity)):
    require(user, 'users_write')
    if any(c.isspace() for c in data.username):
        raise HTTPException(422, 'El usuario no debe contener espacios')
    if data.role not in ROLE_ACTIONS:
        raise HTTPException(422, 'Rol inválido')
    ids = set(data.branch_ids)
    with Session(engine) as db:
        branches = db.scalars(select(Branch).where(Branch.id.in_(ids), Branch.empresa_id == user.empresa_id)).all()
        if len(branches) != len(ids) or len(ids) != len(data.branch_ids):
            raise HTTPException(422, 'Sucursales inválidas o duplicadas')
        created = User(empresa_id=user.empresa_id, username=data.username, password_hash=password_hash.hash(data.password), role=data.role)
        db.add(created)
        try:
            db.flush()
            db.add_all(UserBranch(user_id=created.id, branch_id=branch_id) for branch_id in ids)
            db.add(Audit(empresa_id=user.empresa_id, branch_id=min(ids), action='user_created', record_id=created.id, actor_id=user.id))
            db.commit()
        except Exception:
            db.rollback()
            raise HTTPException(409, 'Usuario duplicado')
        return {'id': created.id, 'username': created.username, 'role': created.role}

def money(v):
    return Decimal(v).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)

def branch_for(db: Session, user: User, branch_id: int):
    branch = db.scalar(select(Branch).where(Branch.id == branch_id, Branch.empresa_id == user.empresa_id))
    if branch is None:
        raise HTTPException(404, 'Sucursal no encontrada para esta empresa')
    if user.role != 'admin_general' and not db.scalar(select(UserBranch).where(UserBranch.user_id == user.id, UserBranch.branch_id == branch_id)):
        raise HTTPException(403, 'Sucursal no autorizada')
    return branch

@app.get('/api/branches')
def branches(user: User = Depends(identity)):
    empresa = user.empresa_id
    with Session(engine) as db:
        return [{'id': b.id, 'name': b.name} for b in db.scalars(select(Branch).where(Branch.empresa_id == empresa).order_by(Branch.id)) if user.role == 'admin_general' or db.scalar(select(UserBranch).where(UserBranch.user_id == user.id, UserBranch.branch_id == b.id))]

def cash_access(db, user, session, owner=False):
    branch_for(db, user, session.branch_id)
    if session.cashier_id != user.id and (owner or user.role not in ('admin_general', 'admin_sucursal', 'contabilidad', 'auditoria')):
        raise HTTPException(403, 'Este turno pertenece a otro cajero')


def cash_totals(db, session):
    sales = db.scalars(select(Sale).where(Sale.cash_session_id == session.id, Sale.payment_method == 'cash')).all()
    movements = db.scalars(select(CashMovement).where(CashMovement.session_id == session.id)).all()
    withdrawn = sum((m.amount for m in movements if m.kind == 'withdrawal'), Decimal('0'))
    deposited = sum((m.amount for m in movements if m.kind == 'deposit'), Decimal('0'))
    refunded = sum((m.amount for m in movements if m.kind == 'refund'), Decimal('0'))
    expected = money(session.opening + sum((s.total for s in sales), Decimal('0')) + deposited - withdrawn - refunded)
    return {'expected': str(expected), 'withdrawn': str(money(withdrawn)), 'deposited': str(money(deposited)), 'cash_sales': len(sales), 'refunded': str(money(refunded))}


def cash_cut(db, session):
    if session.status == 'closed' and session.close_snapshot:
        return json.loads(session.close_snapshot)
    sales = db.scalars(select(Sale).where(Sale.cash_session_id == session.id)).all()
    totals = cash_totals(db, session)
    methods = {method: {'count': sum(s.payment_method == method for s in sales),
        'gross': str(money(sum((s.total for s in sales if s.payment_method == method), Decimal('0'))))}
        for method in ('cash', 'card', 'transfer', 'mercado_pago')}
    # External refunds concern tickets issued in this shift. Cash refunds concern
    # money actually paid out by this shift, including tickets from earlier shifts.
    returns = db.scalars(select(SaleReturn).where(SaleReturn.status == 'completed', SaleReturn.sale_id.in_([s.id for s in sales]))).all() if sales else []
    if session.closed_at:
        cutoff = session.closed_at.replace(tzinfo=None)
        returns = [r for r in returns if r.created_at.replace(tzinfo=None) <= cutoff]
    by_sale = {s.id: s for s in sales}
    for method in ('cash', 'card', 'transfer', 'mercado_pago'):
        methods[method]['ticket_refunds'] = str(money(sum((r.total for r in returns if by_sale[r.sale_id].payment_method == method), Decimal('0'))))
    expected = session.expected_on_close if session.status == 'closed' and session.expected_on_close is not None else Decimal(totals['expected'])
    return {'session_id': session.id, 'sales_count': len(sales), 'payments': methods,
            'gross_sales': str(money(sum((x.total for x in sales), Decimal('0')))),
            'opening': str(session.opening), **totals, 'expected': str(money(expected)),
            'counted': str(session.counted) if session.counted is not None else None,
            'difference': str(money(session.counted - expected)) if session.counted is not None else None,
            'closed_by': session.closed_by, 'closed_at': session.closed_at.isoformat() if session.closed_at else None,
            'snapshot': False}


@app.get('/api/cash/{session_id}/cut')
def get_cash_cut(session_id: int, user: User = Depends(identity)):
    require_any(user, 'cash_open', 'report')
    with Session(engine) as db:
        session = db.scalar(select(CashSession).where(CashSession.id == session_id, CashSession.empresa_id == user.empresa_id))
        if session is None:
            raise HTTPException(404, 'Turno no encontrado')
        cash_access(db, user, session)
        return {**cash_view(db, session), 'cut': cash_cut(db, session)}


def cash_view(db, session):
    cashier = db.get(User, session.cashier_id) if session.cashier_id else None
    register = db.get(CashRegister, session.register_id) if session.register_id else None
    return {'open': session.status == 'open', 'id': session.id, 'register_id': session.register_id,
            'register_name': register.name if register else 'Caja anterior', 'cashier_id': session.cashier_id,
            'cashier_name': cashier.username if cashier else 'Sin responsable (turno anterior)',
            'opening': str(session.opening), 'status': session.status, 'opened_at': session.opened_at.isoformat(),
            'closed_at': session.closed_at.isoformat() if session.closed_at else None,
            'counted': str(session.counted) if session.counted is not None else None,
            **cash_totals(db, session),
            'expected': str(session.expected_on_close) if session.status == 'closed' and session.expected_on_close is not None else cash_totals(db, session)['expected']}

@app.get('/api/cash/registers')
def registers(branch_id: int, user: User = Depends(identity)):
    require_any(user, 'cash_open', 'report')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        return [{'id': r.id, 'name': r.name} for r in db.scalars(select(CashRegister).where(
            CashRegister.branch_id == branch_id, CashRegister.empresa_id == user.empresa_id).order_by(CashRegister.id))]

@app.post('/api/cash/registers', status_code=201)
def create_register(data: RegisterIn, user: User = Depends(identity)):
    require(user, 'cash_deposit')
    if not data.name.strip():
        raise HTTPException(422, 'Nombre de caja vacío')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        register = CashRegister(empresa_id=user.empresa_id, branch_id=data.branch_id, name=data.name.strip())
        db.add(register)
        try:
            db.flush()
            db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, action='register_created', record_id=register.id, actor_id=user.id))
            result = {'id': register.id, 'name': register.name}
            db.commit()
            return result
        except IntegrityError:
            raise HTTPException(409, 'Ya existe una caja con ese nombre')

@app.post('/api/cash/open', status_code=201)
def open_cash(data: OpenCash, user: User = Depends(identity)):
    require(user, 'cash_open')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        db.scalar(select(Branch).where(Branch.id == data.branch_id).with_for_update())
        query = select(CashRegister).where(CashRegister.branch_id == data.branch_id, CashRegister.empresa_id == user.empresa_id)
        if data.register_id:
            query = query.where(CashRegister.id == data.register_id)
        register = db.scalar(query.order_by(CashRegister.id).with_for_update())
        if register is None:
            raise HTTPException(404, 'Caja no encontrada en esta sucursal')
        existing = db.scalar(select(CashSession).where(CashSession.branch_id == data.branch_id, CashSession.status == 'open',
                    or_(CashSession.register_id == register.id, CashSession.cashier_id == user.id)))
        if existing:
            raise HTTPException(409, 'La caja o el cajero ya tiene un turno abierto')
        session = CashSession(empresa_id=user.empresa_id, branch_id=data.branch_id, register_id=register.id,
                              cashier_id=user.id, opening=money(data.opening))
        db.add(session)
        try:
            db.flush()
            db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, action='cash_opened', record_id=session.id, actor_id=user.id))
            result = cash_view(db, session)
            db.commit()
            return result
        except IntegrityError:
            raise HTTPException(409, 'La caja o el cajero ya tiene un turno abierto')

@app.get('/api/cash/current')
def current_cash(branch_id: int, register_id: int | None = None, user: User = Depends(identity)):
    require_any(user, 'cash_open', 'report')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        query = select(CashSession).where(CashSession.empresa_id == user.empresa_id, CashSession.branch_id == branch_id, CashSession.status == 'open')
        if register_id:
            query = query.where(CashSession.register_id == register_id)
        elif user.role not in ('admin_general', 'admin_sucursal', 'contabilidad', 'auditoria'):
            query = query.where(CashSession.cashier_id == user.id)
        session = db.scalar(query.order_by(CashSession.id.desc()))
        if not session:
            return {'open': False}
        cash_access(db, user, session)
        return cash_view(db, session)

@app.get('/api/cash/sessions')
def cash_sessions(branch_id: int, user: User = Depends(identity)):
    require_any(user, 'cash_open', 'report')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        query = select(CashSession).where(CashSession.empresa_id == user.empresa_id, CashSession.branch_id == branch_id)
        if user.role not in ('admin_general', 'admin_sucursal', 'contabilidad', 'auditoria'):
            query = query.where(CashSession.cashier_id == user.id)
        return [cash_view(db, s) for s in db.scalars(query.order_by(CashSession.id.desc()).limit(100))]


def post_cash_movement(session_id, data, user, kind, key):
    require(user, 'cash_withdraw' if kind == 'withdrawal' else 'cash_deposit')
    reason = data.reason.strip()
    if len(reason) < 3 or money(data.amount) <= 0:
        raise HTTPException(422, 'Importe o motivo inválidos')
    with Session(engine) as db:
        session = db.scalar(select(CashSession).where(CashSession.id == session_id, CashSession.empresa_id == user.empresa_id).with_for_update())
        if session is None:
            raise HTTPException(404, 'Turno no encontrado')
        cash_access(db, user, session)
        previous = db.scalar(select(CashMovement).where(CashMovement.empresa_id == user.empresa_id, CashMovement.request_key == key)) if key else None
        if previous:
            if (previous.session_id, previous.amount, previous.kind, previous.reason) != (session.id, money(data.amount), kind, reason):
                raise HTTPException(409, 'Clave de movimiento utilizada con otros datos')
            return {'id': previous.id, 'amount': str(previous.amount), 'replayed': True}
        if session.status != 'open':
            raise HTTPException(409, 'El turno está cerrado')
        available = Decimal(cash_totals(db, session)['expected'])
        amount = money(data.amount)
        if kind == 'withdrawal' and amount > available:
            raise HTTPException(409, 'Efectivo insuficiente en caja')
        movement = CashMovement(empresa_id=user.empresa_id, branch_id=session.branch_id, session_id=session.id,
                               actor_id=user.id, request_key=key, amount=amount, kind=kind, reason=reason)
        db.add(movement)
        try:
            db.flush()
            db.add(Audit(empresa_id=user.empresa_id, branch_id=session.branch_id, action=f'cash_{kind}', record_id=movement.id, actor_id=user.id))
            result = {'id': movement.id, 'amount': str(amount), 'expected_after': str(money(available + (amount if kind == 'deposit' else -amount))), 'replayed': False}
            db.commit()
            return result
        except IntegrityError:
            raise HTTPException(409, 'Movimiento duplicado; vuelve a consultar el turno')

@app.post('/api/cash/{session_id}/withdraw', status_code=201)
def withdraw_cash(session_id: int, data: WithdrawCash, idempotency_key: str | None = Header(default=None, min_length=8, max_length=100), user: User = Depends(identity)):
    return post_cash_movement(session_id, data, user, 'withdrawal', idempotency_key)

@app.post('/api/cash/{session_id}/deposit', status_code=201)
def deposit_cash(session_id: int, data: WithdrawCash, idempotency_key: str | None = Header(default=None, min_length=8, max_length=100), user: User = Depends(identity)):
    return post_cash_movement(session_id, data, user, 'deposit', idempotency_key)

@app.get('/api/cash/{session_id}/movements')
def cash_movements(session_id: int, user: User = Depends(identity)):
    require_any(user, 'cash_open', 'report')
    with Session(engine) as db:
        session = db.scalar(select(CashSession).where(CashSession.id == session_id, CashSession.empresa_id == user.empresa_id))
        if session is None:
            raise HTTPException(404, 'Turno no encontrado')
        cash_access(db, user, session)
        return [{'id': m.id, 'amount': str(m.amount), 'kind': m.kind, 'reason': m.reason,
                 'actor_id': m.actor_id, 'created_at': m.created_at.isoformat(), 'session_id': session.id}
                for m in db.scalars(select(CashMovement).where(CashMovement.session_id == session.id).order_by(CashMovement.id))]

@app.post('/api/cash/{session_id}/close')
def close_cash(session_id: int, data: CloseCash, user: User = Depends(identity)):
    require(user, 'cash_close')
    with Session(engine) as db:
        candidate = db.scalar(select(CashSession).where(CashSession.id == session_id,
            CashSession.empresa_id == user.empresa_id))
        if candidate is None:
            raise HTTPException(404, 'Turno abierto no encontrado')
        cash_access(db, user, candidate)
        db.scalar(select(Branch).where(Branch.id == candidate.branch_id).with_for_update())
        session = db.scalar(select(CashSession).where(CashSession.id == session_id, CashSession.empresa_id == user.empresa_id, CashSession.status == 'open').with_for_update().execution_options(populate_existing=True))
        if session is None:
            raise HTTPException(404, 'Turno abierto no encontrado')
        cash_access(db, user, session)
        expected = Decimal(cash_totals(db, session)['expected'])
        session.counted = money(data.counted)
        session.expected_on_close = expected
        session.closed_by = user.id
        session.closed_at = datetime.now(timezone.utc)
        session.status = 'closed'
        snapshot = cash_cut(db, session)
        snapshot['snapshot'] = True
        session.close_snapshot = json.dumps(snapshot)
        db.add(Audit(empresa_id=user.empresa_id, branch_id=session.branch_id, action='cash_closed', record_id=session.id, actor_id=user.id))
        result = {'id': session.id, 'expected': str(expected), 'counted': str(session.counted), 'difference': str(money(session.counted - expected))}
        db.commit()
        return result

@app.get('/api/ready')
def readiness():
    try:
        with engine.connect() as connection:
            connection.execute(text('SELECT 1'))
    except Exception:
        raise HTTPException(503, 'Base de datos no disponible')
    return {'status': 'ok'}

@app.get('/api/health')
def health():
    return {'status': 'ok', 'mode': os.getenv('APP_ENV', 'demo')}

@app.post('/api/products', status_code=201)
def create_product(data: ProductIn, user: User = Depends(identity)):
    require(user, 'catalog_write')
    empresa = user.empresa_id
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        product = Product(empresa_id=empresa, **data.model_dump(exclude={'branch_id', 'stock', 'initial_cost'}))
        db.add(product)
        try:
            db.flush()
            db.add(Stock(product_id=product.id, branch_id=data.branch_id, quantity=data.stock, average_cost=data.initial_cost))
            db.add(StockMovement(actor_id=user.id, empresa_id=empresa, branch_id=data.branch_id, product_id=product.id, change=data.stock, reason='initial'))
            db.add(Audit(empresa_id=empresa, branch_id=data.branch_id, action='product_created', record_id=product.id, actor_id=user.id))
            db.commit()
        except Exception:
            db.rollback()
            raise HTTPException(409, 'SKU duplicado o datos inválidos')
        return {'id': product.id, 'sku': product.sku}

@app.get('/api/products')
def products(branch_id: int, include_unstocked: bool = False, user: User = Depends(identity)):
    empresa = user.empresa_id
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        query = select(Product, Stock).outerjoin(Stock, and_(Stock.product_id == Product.id, Stock.branch_id == branch_id)).where(Product.empresa_id == empresa)
        if not include_unstocked:
            query = query.where(Stock.id.is_not(None))
        rows = db.execute(query.order_by(Product.name)).all()
        return [product_view(p, s) for p, s in rows]

@app.post('/api/customers', status_code=201)
def create_customer(data: CustomerIn, user: User = Depends(identity)):
    require(user, 'customer_write')
    name = data.name.strip()
    phone = data.phone.strip() if data.phone else None
    if len(name) < 2:
        raise HTTPException(422, 'Nombre inválido')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        customer = Customer(empresa_id=user.empresa_id, name=name, phone=phone or None)
        db.add(customer)
        db.flush()
        db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id,
                     action='customer_created', record_id=customer.id, actor_id=user.id))
        result = {'id': customer.id, 'name': customer.name, 'phone': customer.phone}
        db.commit()
        return result

@app.get('/api/customers')
def customers(branch_id: int, q: str = '', user: User = Depends(identity)):
    require(user, 'customer_read')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        query = select(Customer).where(Customer.empresa_id == user.empresa_id)
        if q.strip():
            term = f'%{q.strip()[:80]}%'
            query = query.where(or_(Customer.name.ilike(term), Customer.phone.ilike(term)))
        rows = db.scalars(query.order_by(Customer.id.desc()).limit(50)).all()
        return [{'id': x.id, 'name': x.name, 'phone': x.phone} for x in rows]

@app.put('/api/customers/{customer_id}')
def update_customer(customer_id: int, data: CustomerUpdate, user: User = Depends(identity)):
    require(user, 'customer_write')
    name = data.name.strip()
    if len(name) < 2:
        raise HTTPException(422, 'Nombre inválido')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        customer = db.scalar(select(Customer).where(Customer.id == customer_id,
                                                   Customer.empresa_id == user.empresa_id))
        if customer is None:
            raise HTTPException(404, 'Cliente no encontrado')
        customer.name = name
        customer.phone = (data.phone or '').strip() or None
        db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id,
                     action='customer_updated', record_id=customer.id, actor_id=user.id))
        result = {'id': customer.id, 'name': customer.name, 'phone': customer.phone}
        db.commit()
        return result

@app.get('/api/customers/{customer_id}/sales')
def customer_sales(customer_id: int, branch_id: int, user: User = Depends(identity)):
    require(user, 'customer_read')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        if not db.scalar(select(Customer.id).where(Customer.id == customer_id,
                                                  Customer.empresa_id == user.empresa_id)):
            raise HTTPException(404, 'Cliente no encontrado')
        rows = db.scalars(select(Sale).where(Sale.customer_id == customer_id,
                          Sale.empresa_id == user.empresa_id, Sale.branch_id == branch_id)
                          .order_by(Sale.id.desc()).limit(50)).all()
        return [{'id': s.id, 'created_at': s.created_at.isoformat(), 'total': str(s.total),
                 'payment_method': s.payment_method} for s in rows]

@app.post('/api/products/{product_id}/stock')
def adjust_stock(product_id: int, data: AdjustStock, user: User = Depends(identity)):
    require(user, 'stock_write')
    empresa = user.empresa_id
    if data.change == 0:
        raise HTTPException(422, 'El movimiento debe cambiar existencias')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        db.scalar(select(Branch).where(Branch.id == data.branch_id).with_for_update())
        row = db.execute(select(Product, Stock).join(Stock).where(Product.id == product_id, Product.empresa_id == empresa, Stock.branch_id == data.branch_id).with_for_update()).first()
        if not row:
            raise HTTPException(404, 'Producto sin inventario en sucursal')
        product, stock = row
        if stock.quantity + data.change < reserved_quantity(db, data.branch_id, product_id):
            raise HTTPException(409, 'Inventario insuficiente')
        stock.quantity += data.change
        db.add(StockMovement(actor_id=user.id, empresa_id=empresa, branch_id=data.branch_id, product_id=product.id, change=data.change, reason=data.reason))
        db.add(Audit(empresa_id=empresa, branch_id=data.branch_id, action='stock_adjusted', record_id=product.id, actor_id=user.id))
        result = {'product_id': product.id, 'stock': stock.quantity}
        db.commit()
        return result

@app.get('/api/stock/movements')
def movements(branch_id: int, user: User = Depends(identity)):
    require(user, 'stock_write')
    empresa = user.empresa_id
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        rows = db.scalars(select(StockMovement).where(StockMovement.empresa_id == empresa, StockMovement.branch_id == branch_id).order_by(StockMovement.id.desc()).limit(100)).all()
        return [{'id': x.id, 'product_id': x.product_id, 'change': x.change, 'reason': x.reason, 'reference_id': x.reference_id, 'actor_id': x.actor_id, 'created_at': x.created_at.isoformat()} for x in rows]

@app.post('/api/stock/transfers', status_code=201)
def transfer_stock(data: TransferStock, idempotency_key: str = Header(min_length=8, max_length=100), user: User = Depends(identity)):
    require(user, 'stock_write')
    if data.source_branch_id == data.target_branch_id:
        raise HTTPException(422, 'Selecciona dos sucursales distintas')
    with Session(engine) as db:
        try:
            branch_for(db, user, data.source_branch_id)
            branch_for(db, user, data.target_branch_id)
            # Consistent lock order serializes opposite-direction transfers in PostgreSQL.
            db.scalars(select(Branch).where(Branch.id.in_((data.source_branch_id, data.target_branch_id)))
                       .order_by(Branch.id).with_for_update()).all()
            existing = db.scalar(select(StockTransfer).where(
                StockTransfer.empresa_id == user.empresa_id, StockTransfer.request_key == idempotency_key))
            if existing:
                if (existing.source_branch_id, existing.target_branch_id, existing.product_id, existing.quantity) != (
                    data.source_branch_id, data.target_branch_id, data.product_id, data.quantity):
                    raise HTTPException(409, 'Clave de traspaso ya utilizada para otra operación')
                return {'id': existing.id, 'source_branch_id': existing.source_branch_id,
                        'target_branch_id': existing.target_branch_id, 'product_id': existing.product_id,
                        'quantity': existing.quantity, 'replayed': True}
            product = db.scalar(select(Product).where(Product.id == data.product_id, Product.empresa_id == user.empresa_id))
            if product is None:
                raise HTTPException(404, 'Producto no encontrado para esta empresa')
            stocks = db.scalars(select(Stock).where(Stock.product_id == product.id,
                Stock.branch_id.in_((data.source_branch_id, data.target_branch_id)))
                .order_by(Stock.branch_id).with_for_update()).all()
            by_branch = {stock.branch_id: stock for stock in stocks}
            source = by_branch.get(data.source_branch_id)
            if source is None or source.quantity - reserved_quantity(db, data.source_branch_id, data.product_id) < data.quantity:
                raise HTTPException(409, 'Existencias insuficientes en sucursal origen')
            target = by_branch.get(data.target_branch_id)
            if target is None:
                target = Stock(product_id=product.id, branch_id=data.target_branch_id, quantity=0)
                db.add(target)
            target.average_cost = weighted_cost(target.quantity, target.average_cost or Decimal('0'), data.quantity, source.average_cost or Decimal('0'))
            source.quantity -= data.quantity
            target.quantity += data.quantity
            transfer = StockTransfer(empresa_id=user.empresa_id, source_branch_id=data.source_branch_id,
                target_branch_id=data.target_branch_id, product_id=product.id, quantity=data.quantity,
                request_key=idempotency_key, actor_id=user.id)
            db.add(transfer)
            db.flush()
            db.add_all([
                StockMovement(actor_id=user.id, empresa_id=user.empresa_id, branch_id=data.source_branch_id,
                    product_id=product.id, change=-data.quantity, reason='transfer_out', reference_id=transfer.id),
                StockMovement(actor_id=user.id, empresa_id=user.empresa_id, branch_id=data.target_branch_id,
                    product_id=product.id, change=data.quantity, reason='transfer_in', reference_id=transfer.id),
                Audit(empresa_id=user.empresa_id, branch_id=data.source_branch_id,
                    action='stock_transferred', record_id=transfer.id, actor_id=user.id),
            ])
            result = {'id': transfer.id, 'source_branch_id': data.source_branch_id,
                      'target_branch_id': data.target_branch_id, 'product_id': product.id,
                      'quantity': data.quantity, 'replayed': False}
            db.commit()
            return result
        except IntegrityError:
            db.rollback()
            raise HTTPException(409, 'Traspaso duplicado o inventario modificado simultáneamente')
        except Exception:
            db.rollback()
            raise

def product_view(p, stock):
    return {'id': p.id, 'sku': p.sku, 'barcode': p.barcode, 'name': p.name, 'price': str(p.price),
            'unit': p.unit, 'tax_rate': str(p.tax_rate), 'tax_exempt': p.tax_exempt,
            'price_includes_tax': p.price_includes_tax, 'active': p.active, 'stock': stock.quantity if stock else 0,
            'minimum': stock.minimum if stock else 0, 'low_stock': bool(p.active and stock and stock.minimum > 0 and stock.quantity <= stock.minimum),
            'average_cost': str(stock.average_cost) if stock else '0.0000'}


def weighted_cost(quantity, cost, incoming, incoming_cost):
    return ((quantity * cost + incoming * incoming_cost) / (quantity + incoming)).quantize(Decimal('.0001'), rounding=ROUND_HALF_UP)

@app.put('/api/products/{product_id}')
def edit_product(product_id: int, data: ProductUpdate, user: User = Depends(identity)):
    require(user, 'catalog_write')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        product = db.scalar(select(Product).where(Product.id == product_id, Product.empresa_id == user.empresa_id).with_for_update())
        if product is None:
            raise HTTPException(404, 'Producto no encontrado')
        for key, value in data.model_dump(exclude={'branch_id'}).items():
            setattr(product, key, value)
        db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, action='product_updated', record_id=product.id, actor_id=user.id))
        try:
            db.commit()
            return {'id': product.id, 'sku': product.sku}
        except IntegrityError:
            raise HTTPException(409, 'SKU o código de barras duplicado')


def approve_discount(db, user, data):
    if not data.discount_percent:
        return None
    if not data.discount_reason or len(data.discount_reason.strip()) < 3:
        raise HTTPException(422, 'Indica el motivo del descuento')
    approver = user
    if data.approval:
        approver = db.scalar(select(User).where(User.username == data.approval.username, User.empresa_id == user.empresa_id))
        if approver is None or not approver.active or not password_hash.verify(data.approval.password, approver.password_hash):
            raise HTTPException(403, 'Autorización de descuento inválida')
    require(approver, 'discount')
    branch_for(db, approver, data.branch_id)
    return approver.id


def price_line(product, quantity, percent):
    # Round each line, so printed totals equal the sum of the stored lines.
    gross = money(product.price * quantity)
    discount = money(gross * percent / Decimal('100'))
    discounted = gross - discount
    rate = Decimal('0') if product.tax_exempt else product.tax_rate
    net = money(discounted / (Decimal('1') + rate)) if product.price_includes_tax else discounted
    tax = discounted - net if product.price_includes_tax else money(net * rate)
    return net, tax, discount


@app.post('/api/sales/quote')
def sale_quote(data: SaleIn, user: User = Depends(identity)):
    require(user, 'sale')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        ids = [x.product_id for x in data.items]
        rows = db.scalars(select(Product).where(Product.empresa_id == user.empresa_id, Product.id.in_(ids), Product.active == True)).all()
        if len(rows) != len(ids) or len(ids) != len(set(ids)):
            raise HTTPException(422, 'Carrito con productos inválidos o duplicados')
        by_id = {p.id: p for p in rows}
        amounts = [price_line(by_id[x.product_id], x.quantity, data.discount_percent) for x in data.items]
        subtotal = sum((a[0] for a in amounts), Decimal('0'))
        tax = sum((a[1] for a in amounts), Decimal('0'))
        return {'subtotal': str(subtotal), 'tax': str(tax), 'discount_total': str(sum((a[2] for a in amounts), Decimal('0'))), 'total': str(subtotal + tax)}


def sale_result(sale, replayed=False):
    return {'id': sale.id, 'folio': sale.folio or f'LI-B{sale.branch_id}-{sale.id:08d}',
            'subtotal': str(sale.subtotal), 'tax': str(sale.tax), 'discount_total': str(sale.discount_total),
            'total': str(sale.total), 'change': str(money(sale.paid - sale.total)), 'replayed': replayed}

@app.post('/api/sales', status_code=201)
def sell(data: SaleIn, idempotency_key: str = Header(min_length=8, max_length=100), user: User = Depends(identity)):
    require(user, 'sale')
    if data.payment_method not in ('cash', 'card', 'transfer', 'mercado_pago'):
        raise HTTPException(422, 'Forma de pago inválida')
    ids = [i.product_id for i in data.items]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, 'Productos duplicados en carrito')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        db.scalar(select(Branch).where(Branch.id == data.branch_id).with_for_update())
        existing = db.scalar(select(Sale).where(Sale.empresa_id == user.empresa_id, Sale.request_key == idempotency_key))
        if existing:
            previous = sorted((x.product_id, x.quantity) for x in existing.items)
            incoming = sorted((x.product_id, x.quantity) for x in data.items)
            if (existing.branch_id != data.branch_id or existing.customer_id != data.customer_id
                or existing.payment_method != data.payment_method or existing.paid != money(data.paid)
                or previous != incoming or existing.discount_percent != data.discount_percent
                or (existing.discount_reason or '') != (data.discount_reason or '').strip()
                or (data.cash_session_id is not None and existing.cash_session_id != data.cash_session_id)):
                raise HTTPException(409, 'Clave de venta ya utilizada para otra operación')
            if existing.actor_id is not None and existing.actor_id != user.id:
                raise HTTPException(403, 'El reintento pertenece a otro cajero')
            return sale_result(existing, True)
        intent = None
        if data.payment_method == 'mercado_pago':
            intent = db.scalar(select(PaymentIntent).where(PaymentIntent.id == data.payment_intent_id,
                PaymentIntent.empresa_id == user.empresa_id, PaymentIntent.actor_id == user.id).with_for_update())
            if intent is None or intent.status != 'approved' or intent.review_reason or intent.cancel_requested_at:
                raise HTTPException(409, 'El pago integrado no está confirmado')
            original = json.loads(intent.payload)
            comparable = data.model_dump(mode='json', exclude={'approval','payment_intent_id'})
            if comparable.get('cash_session_id') != original.get('cash_session_id'):
                previous_cash = db.scalar(select(CashSession).where(
                    CashSession.id == original.get('cash_session_id'),
                    CashSession.empresa_id == user.empresa_id,
                    CashSession.branch_id == intent.branch_id,
                    CashSession.cashier_id == user.id))
                if previous_cash is None or previous_cash.status != 'closed':
                    raise HTTPException(409, 'Solo puedes cambiar el turno de entrega cuando el turno original ya está cerrado')
                comparable['cash_session_id'] = original.get('cash_session_id')
            if comparable != original:
                raise HTTPException(409, 'El carrito no coincide con el pago integrado')
        elif data.payment_intent_id:
            raise HTTPException(422, 'Referencia de pago integrada con método inválido')
        customer = db.scalar(select(Customer).where(Customer.id == data.customer_id, Customer.empresa_id == user.empresa_id)) if data.customer_id else None
        if data.customer_id and customer is None:
            raise HTTPException(404, 'Cliente fuera de esta empresa o inexistente')
        query = select(CashSession).where(CashSession.empresa_id == user.empresa_id, CashSession.branch_id == data.branch_id, CashSession.status == 'open')
        query = query.where(CashSession.id == data.cash_session_id) if data.cash_session_id else query.where(CashSession.cashier_id == user.id)
        cash = db.scalar(query.with_for_update())
        if not cash:
            raise HTTPException(409, 'Abre tu turno de caja antes de vender')
        cash_access(db, user, cash, owner=True)
        approved_by = approve_discount(db, user, data)
        rows = db.execute(select(Product, Stock).join(Stock).where(Product.id.in_(ids), Product.empresa_id == user.empresa_id,
                        Stock.branch_id == data.branch_id).order_by(Product.id).with_for_update()).all()
        mapped = {p.id: (p, stock) for p, stock in rows}
        if len(mapped) != len(ids):
            raise HTTPException(404, 'Producto fuera de esta empresa o sucursal')
        subtotal, tax, discount = Decimal('0'), Decimal('0'), Decimal('0')
        sale_lines = []
        for item in data.items:
            product, stock = mapped[item.product_id]
            if intent and intent.price_snapshot:
                from types import SimpleNamespace
                values = json.loads(intent.price_snapshot)[str(product.id)]
                product = SimpleNamespace(**values)
                product.price = Decimal(product.price)
                product.tax_rate = Decimal(product.tax_rate)
            elif not product.active:
                raise HTTPException(409, f'Producto inactivo: {product.name}')
            if stock.quantity - reserved_quantity(db, data.branch_id, item.product_id, intent.id if intent else None) < item.quantity:
                raise HTTPException(409, f'Sin existencias suficientes: {product.name}')
            net, line_tax, line_discount = price_line(product, item.quantity, data.discount_percent)
            subtotal += net
            tax += line_tax
            discount += line_discount
            stock.quantity -= item.quantity
            sale_lines.append(SaleItem(product_id=product.id, name=product.name, quantity=item.quantity, unit_price=product.price,
                unit=product.unit, tax_rate=product.tax_rate, tax_exempt=product.tax_exempt, price_includes_tax=product.price_includes_tax,
                discount_amount=line_discount, net_amount=net, tax_amount=line_tax, cost=stock.average_cost))
        total = subtotal + tax
        if total > Decimal('9999999999.99') or subtotal > Decimal('9999999999.99') or discount > Decimal('9999999999.99'):
            raise HTTPException(422, 'El importe excede el límite de una venta')
        if intent and money(intent.amount) != total:
            raise HTTPException(409, 'El precio cambió después del pago; requiere conciliación')
        if data.payment_method == 'cash' and money(data.paid) < total:
            raise HTTPException(422, 'Pago insuficiente')
        if data.payment_method != 'cash' and money(data.paid) != total:
            raise HTTPException(422, 'El pago debe coincidir con el total')
        sale = Sale(empresa_id=user.empresa_id, branch_id=data.branch_id, customer_id=data.customer_id,
                    customer_name=customer.name if customer else None, cash_session_id=cash.id, actor_id=user.id,
                    discount_percent=data.discount_percent, discount_total=discount, discount_reason=(data.discount_reason or '').strip() or None,
                    discount_approved_by=approved_by, request_key=idempotency_key, subtotal=subtotal, tax=tax,
                    total=total, payment_method=data.payment_method, paid=money(data.paid), items=sale_lines)
        db.add(sale)
        try:
            db.flush()
            sale.folio = f'LI-B{data.branch_id}-{sale.id:08d}'
            for item in data.items:
                db.add(StockMovement(empresa_id=user.empresa_id, branch_id=data.branch_id, product_id=item.product_id,
                       change=-item.quantity, reason='sale', reference_id=sale.id, actor_id=user.id))
            db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, action='sale_created', record_id=sale.id, actor_id=user.id))
            if approved_by:
                db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, action='discount_approved', record_id=sale.id, actor_id=approved_by))
            if intent:
                intent.status = 'completed'
                intent.reservation_active = False
                intent.delivery_cash_session_id = cash.id
                intent.delivered_at = datetime.now(timezone.utc)
                if cash.id != json.loads(intent.payload).get('cash_session_id'):
                    db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id,
                        action='payment_delivered_new_shift', record_id=sale.id, actor_id=user.id))
            result = sale_result(sale)
            db.commit()
            return result
        except IntegrityError:
            raise HTTPException(409, 'Venta duplicada; vuelve a consultar el historial')

@app.get('/api/sales')
def sales(branch_id: int, user: User = Depends(identity)):
    require_any(user, 'sale', 'report')
    empresa = user.empresa_id
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        rows = db.scalars(select(Sale).where(Sale.empresa_id == empresa, Sale.branch_id == branch_id).order_by(Sale.id.desc()).limit(30)).all()
        return [{'id': s.id, 'folio': s.folio, 'status': 'cancelled' if db.scalar(select(SaleReturn.id).where(SaleReturn.sale_id == s.id, SaleReturn.kind == 'cancellation', SaleReturn.status == 'completed')) else 'completed', 'total': str(s.total), 'created_at': s.created_at.isoformat(), 'payment_method': s.payment_method} for s in rows]

@app.get('/api/sales/{sale_id}')
def sale_detail(sale_id: int, user: User = Depends(identity)):
    require_any(user, 'sale', 'report')
    with Session(engine) as db:
        sale = db.scalar(select(Sale).where(Sale.id == sale_id, Sale.empresa_id == user.empresa_id))
        if sale is None:
            raise HTTPException(404, 'Venta no encontrada')
        branch = branch_for(db, user, sale.branch_id)
        customer = db.get(Customer, sale.customer_id) if sale.customer_id else None
        return {
            'folio': sale.folio, 'discount_total': str(sale.discount_total), 'discount_percent': str(sale.discount_percent),
            'discount_reason': sale.discount_reason, 'discount_approved_by': sale.discount_approved_by,
            'cash_session_id': sale.cash_session_id, 'actor_id': sale.actor_id,
            'status': 'cancelled' if db.scalar(select(SaleReturn.id).where(SaleReturn.sale_id == sale.id, SaleReturn.kind == 'cancellation', SaleReturn.status == 'completed')) else 'completed',
            'id': sale.id, 'branch_id': branch.id, 'branch_name': branch.name,
            'customer_id': sale.customer_id, 'customer_name': sale.customer_name or (customer.name if customer else None),
            'created_at': sale.created_at.isoformat(), 'payment_method': sale.payment_method,
            'subtotal': str(sale.subtotal), 'tax': str(sale.tax), 'total': str(sale.total),
            'paid': str(sale.paid), 'change': str(money(sale.paid - sale.total)),
            'items': [
                {'id': item.id, 'product_id': item.product_id, 'name': item.name, 'quantity': item.quantity,
                 'returned_quantity': returned_quantities(db, sale.id).get(item.id, 0),
                 'refundable_total': str(refundable_line_total(sale, item)) if item.net_amount is not None or not sale.discount_total else None,
                 'unit_price': str(item.unit_price), 'unit': item.unit, 'tax_rate': str(item.tax_rate), 'tax_exempt': item.tax_exempt,
                 'discount': str(item.discount_amount), 'net': str(item.net_amount) if item.net_amount is not None else None,
                 'tax': str(item.tax_amount) if item.tax_amount is not None else None,
                 'line_total': str(item.net_amount) if item.net_amount is not None else str(money(item.unit_price * item.quantity))}
                for item in sale.items
            ],
        }

class ReturnLine(BaseModel):
    sale_item_id: int = Field(gt=0)
    quantity: int = Field(gt=0, le=10000)
    restock: bool = True

class ReturnIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    reason: str = Field(min_length=3, max_length=160)
    cash_session_id: int | None = Field(default=None, gt=0)
    payment_reference: str | None = Field(default=None, max_length=100)
    items: list[ReturnLine] = Field(min_length=1, max_length=100)


def returned_quantities(db, sale_id):
    rows = db.scalars(select(SaleReturnItem).join(SaleReturn).where(SaleReturn.sale_id == sale_id, SaleReturn.status == 'completed')).all()
    totals = {}
    for item in rows:
        totals[item.sale_item_id] = totals.get(item.sale_item_id, 0) + item.quantity
    return totals


def return_view(record, replayed=False):
    return {'id': record.id, 'kind': record.kind, 'sale_id': record.sale_id, 'total': str(record.total),
            'status':record.status,'provider_refund_id':record.provider_refund_id,
            'requested_at':record.requested_at.isoformat() if record.requested_at else None,
            'reason': record.reason, 'actor_id': record.actor_id, 'cash_session_id': record.cash_session_id,
            'payment_reference': record.payment_reference, 'created_at': record.created_at.isoformat(),
            'items': [{'sale_item_id': x.sale_item_id, 'quantity': x.quantity, 'restock': x.restock,
                       'total': str(x.total)} for x in record.items], 'replayed': replayed}


def refundable_line_total(sale, item):
    if item.net_amount is not None and item.tax_amount is not None:
        return item.net_amount + item.tax_amount
    # Legacy tickets lack stored line taxes. Allocate the original total; refuse
    # discounted legacy records whose allocation cannot be established reliably.
    if sale.discount_total:
        raise HTTPException(409, 'Venta antigua sin desglose: requiere revisión de administración')
    gross = sum((x.unit_price * x.quantity for x in sale.items), Decimal('0'))
    if gross <= 0:
        raise HTTPException(409, 'Venta antigua sin importes válidos')
    ordered = sorted(sale.items, key=lambda x: x.id)
    before = sum((x.unit_price * x.quantity for x in ordered if x.id < item.id), Decimal('0'))
    return money(sale.total * (before + item.unit_price * item.quantity) / gross) - money(sale.total * before / gross)


@app.get('/api/sales/{sale_id}/returns')
def sale_returns(sale_id: int, user: User = Depends(identity)):
    require_any(user, 'sale', 'report', 'sale_return')
    with Session(engine) as db:
        sale = db.scalar(select(Sale).where(Sale.id == sale_id, Sale.empresa_id == user.empresa_id))
        if sale is None:
            raise HTTPException(404, 'Venta no encontrada')
        branch_for(db, user, sale.branch_id)
        return [return_view(r) for r in db.scalars(select(SaleReturn).where(SaleReturn.sale_id == sale.id).order_by(SaleReturn.id.desc()))]


@app.post('/api/sales/{sale_id}/returns', status_code=201)
def create_return(sale_id: int, data: ReturnIn, idempotency_key: str = Header(min_length=8, max_length=100), user: User = Depends(identity)):
    return process_return(sale_id, data, idempotency_key, user)


def prepare_return(sale_id, data, idempotency_key, user, cancellation=False):
    require(user, 'sale_cancel' if cancellation else 'sale_return')
    ids = [x.sale_item_id for x in data.items]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, 'Partidas duplicadas')
    payload = json.dumps({'sale_id': sale_id, 'kind': 'cancellation' if cancellation else 'return', **data.model_dump(mode='json'),
                          'items': sorted([x.model_dump() for x in data.items], key=lambda x: x['sale_item_id'])}, sort_keys=True)
    with Session(engine) as db:
        sale = db.scalar(select(Sale).where(Sale.id == sale_id, Sale.empresa_id == user.empresa_id))
        if sale is None:
            raise HTTPException(404, 'Venta no encontrada')
        branch_for(db, user, sale.branch_id)
        # Same lock order as checkout: branch, cash, stock. Serializes returns
        # against each other and against checkout in this branch on PostgreSQL.
        db.scalar(select(Branch).where(Branch.id == sale.branch_id).with_for_update())
        existing = db.scalar(select(SaleReturn).where(SaleReturn.empresa_id == user.empresa_id, SaleReturn.request_key == idempotency_key))
        if existing:
            if existing.actor_id != user.id or (json.loads(existing.payload) | {'kind': existing.kind}) != json.loads(payload):
                raise HTTPException(409, 'Clave de devolución utilizada para otra operación')
            return return_view(existing, True)
        if db.scalar(select(SaleReturn.id).where(SaleReturn.sale_id == sale.id, SaleReturn.status != 'completed')):
            raise HTTPException(409, 'Hay una devolución pendiente de confirmar; reintenta esa operación')
        if db.scalar(select(SaleReturn.id).where(SaleReturn.sale_id == sale.id, SaleReturn.kind == 'cancellation')):
            raise HTTPException(409, 'La venta ya está cancelada')
        if cancellation and db.scalar(select(SaleReturn.id).where(SaleReturn.sale_id == sale.id)):
            raise HTTPException(409, 'La venta tiene devoluciones; completa la devolución de las unidades pendientes')
        cash = None
        if sale.payment_method == 'cash':
            cash = db.scalar(select(CashSession).where(CashSession.id == data.cash_session_id,
                CashSession.empresa_id == user.empresa_id, CashSession.branch_id == sale.branch_id).with_for_update())
            if cash is None or cash.status != 'open':
                raise HTTPException(409, 'Selecciona un turno abierto de la misma sucursal para reembolsar')
        elif sale.payment_method == 'mercado_pago':
            if data.cash_session_id is not None or data.payment_reference:
                raise HTTPException(422, 'La devolución integrada no utiliza efectivo ni referencia manual')
        elif not data.payment_reference:
            raise HTTPException(422, 'Indica la referencia del reembolso externo realizado')
        elif data.cash_session_id is not None:
            raise HTTPException(422, 'Un reembolso externo no debe afectar efectivo')
        already = returned_quantities(db, sale.id)
        by_id = {x.id: x for x in sale.items}
        lines, total = [], Decimal('0')
        for incoming in data.items:
            item = by_id.get(incoming.sale_item_id)
            if item is None:
                raise HTTPException(422, 'Partida fuera de esta venta')
            previous = already.get(item.id, 0)
            if previous + incoming.quantity > item.quantity:
                raise HTTPException(409, 'La cantidad supera lo pendiente por devolver')
            original = refundable_line_total(sale, item)
            # Cumulative rounding leaves no lost/extra cents over partial returns.
            amount = money(original * (previous + incoming.quantity) / item.quantity) - money(original * previous / item.quantity)
            total += amount
            lines.append(SaleReturnItem(sale_item_id=item.id, quantity=incoming.quantity, restock=incoming.restock, total=amount))
        if cash and Decimal(cash_totals(db, cash)['expected']) < total:
            raise HTTPException(409, 'Efectivo insuficiente en el turno para el reembolso')
        record = SaleReturn(kind='cancellation' if cancellation else 'return', empresa_id=user.empresa_id, sale_id=sale.id, actor_id=user.id,
            cash_session_id=cash.id if cash else None, reason=data.reason, payment_reference=data.payment_reference,
            payload=payload, total=total, request_key=idempotency_key, items=lines)
        db.add(record)
        db.flush()
        if sale.payment_method == 'mercado_pago':
            config=mp.settings()
            intent=db.scalar(select(PaymentIntent).where(PaymentIntent.empresa_id==user.empresa_id,
                PaymentIntent.id==sale.request_key.removeprefix('mp-sale-')))
            if (intent is None or intent.status!='completed' or not intent.payment_id
                or intent.empresa_id!=config['company'] or intent.amount!=sale.total):
                raise HTTPException(409,'No se identifica el pago integrado del ticket')
            if total<=0:
                raise HTTPException(409,'La devolución integrada debe tener un importe positivo')
            import uuid
            record.status='pending';record.requested_at=datetime.now(timezone.utc)
            record.provider_key=str(uuid.uuid4());record.provider_payment_id=intent.payment_id
            record.provider_refunded_before=sum((r.total for r in db.scalars(select(SaleReturn).where(
                SaleReturn.sale_id==sale.id,SaleReturn.status=='completed'))),Decimal(0))
            # Stock rows must exist before any external refund is requested.
            for incoming in data.items:
                if incoming.restock and not db.scalar(select(Stock.id).where(Stock.product_id==by_id[incoming.sale_item_id].product_id,Stock.branch_id==sale.branch_id)):
                    raise HTTPException(409,'Existencia no encontrada; revisa el inventario antes de reembolsar')
            db.add(Audit(empresa_id=user.empresa_id,branch_id=sale.branch_id,actor_id=user.id,
                action='integrated_return_requested',record_id=record.id))
            db.commit()
            return return_view(record)
        return apply_return_record(db,sale,record,data,user,cash)


def apply_return_record(db, sale, record, data, user, cash=None):
    by_id={x.id:x for x in sale.items}
    total=record.total
    cancellation=record.kind=='cancellation'
    if record.status!='completed':
        record.created_at=datetime.now(timezone.utc)
    record.status='completed';record.finalized_at=datetime.now(timezone.utc)
    for incoming in sorted(data.items, key=lambda x: by_id[x.sale_item_id].product_id):
        if not incoming.restock:
            continue
        item = by_id[incoming.sale_item_id]
        stock = db.scalar(select(Stock).where(Stock.product_id == item.product_id, Stock.branch_id == sale.branch_id).with_for_update())
        if stock is None:
            raise HTTPException(409, 'Existencia de la venta no encontrada; revisa el inventario')
        new_quantity = stock.quantity + incoming.quantity
        stock.average_cost = ((stock.average_cost * stock.quantity + item.cost * incoming.quantity) / new_quantity).quantize(Decimal('.0001'), rounding=ROUND_HALF_UP)
        stock.quantity = new_quantity
        db.add(StockMovement(empresa_id=user.empresa_id, branch_id=sale.branch_id, product_id=item.product_id,
            actor_id=user.id, change=incoming.quantity, reason='sale_cancel' if cancellation else 'sale_return', reference_id=record.id))
    if cash:
        db.add(CashMovement(empresa_id=user.empresa_id, branch_id=sale.branch_id, session_id=cash.id,
            actor_id=user.id, amount=total, kind='refund', reason=f'Devolución #{record.id}: {data.reason}'[:160]))
    db.add(Audit(empresa_id=user.empresa_id, branch_id=sale.branch_id, actor_id=user.id, action='sale_cancel' if cancellation else 'sale_return', record_id=record.id))
    db.flush()
    result = return_view(record)
    db.commit()
    return result


def process_return(sale_id, data, idempotency_key, user, cancellation=False):
    result=prepare_return(sale_id,data,idempotency_key,user,cancellation)
    if result['status']=='completed':return result
    return complete_integrated_return(result['id'],user)


def complete_integrated_return(return_id, user):
    config=mp.settings()
    with Session(engine) as db:
        record=db.scalar(select(SaleReturn).where(SaleReturn.id==return_id,SaleReturn.empresa_id==user.empresa_id))
        if record is None:raise HTTPException(404,'Devolución no encontrada')
        require(user,'sale_cancel' if record.kind=='cancellation' else 'sale_return')
        sale=db.get(Sale,record.sale_id);branch_for(db,user,sale.branch_id)
        if record.status=='completed':return return_view(record,True)
        intent=db.get(PaymentIntent,sale.request_key.removeprefix('mp-sale-'))
        if (intent is None or intent.payment_id!=record.provider_payment_id or intent.empresa_id!=config['company']):
            raise HTTPException(409,'No se identifica el pago del ticket')
        db.expunge(intent)
        amount=record.total;baseline=record.provider_refunded_before;key=record.provider_key
        payment_id=record.provider_payment_id;branch_id=sale.branch_id
    payment=mp.call('GET','/v1/payments/'+payment_id)
    validate_cancel_payment(payment,intent,config)
    if str(payment.get('id'))!=payment_id or payment.get('status') not in ('approved','refunded'):
        raise HTTPException(409,'Pago no reembolsable; operación pendiente')
    refunded=Decimal(str(payment.get('transaction_amount_refunded',0)))
    target=baseline+amount
    if refunded==baseline:
        result=mp.call('POST','/v1/payments/'+payment_id+'/refunds',{'amount':float(amount)},key=key)
        if str(result.get('payment_id'))!=payment_id or Decimal(str(result.get('amount',0)))!=amount:
            raise HTTPException(502,'Respuesta de reembolso inválida; operación pendiente')
        with Session(engine) as db:
            stored=db.get(SaleReturn,return_id)
            stored.provider_refund_id=str(result.get('id',''))[:100] or None
            db.commit()
        payment=mp.call('GET','/v1/payments/'+payment_id)
        validate_cancel_payment(payment,intent,config)
        refunded=Decimal(str(payment.get('transaction_amount_refunded',0)))
    if refunded!=target or payment.get('status') not in ('approved','refunded'):
        raise HTTPException(409,'El reembolso del proveedor no coincide con la devolución pendiente; conserva la operación y vuelve a consultar')
    with Session(engine) as db:
        db.scalar(select(Branch).where(Branch.id==branch_id).with_for_update())
        record=db.scalar(select(SaleReturn).where(SaleReturn.id==return_id).with_for_update().execution_options(populate_existing=True))
        if record.status=='completed':return return_view(record,True)
        sale=db.get(Sale,record.sale_id)
        completed=sum((r.total for r in db.scalars(select(SaleReturn).where(SaleReturn.sale_id==sale.id,SaleReturn.status=='completed'))),Decimal(0))
        if completed!=baseline:
            raise HTTPException(409,'Cambió la conciliación local; operación pendiente')
        current=db.get(PaymentIntent,intent.id)
        record.payment_reference='MP:'+ (record.provider_refund_id or ('verificado-'+str(record.id)))[:97]
        record_payment_observation(db,current,payment)
        # A verified primary refund remains reconcilable against the completed
        # return ledger; never clear unrelated duplicate/chargeback alerts here.
        data=ReturnIn(**json.loads(record.payload))
        return apply_return_record(db,sale,record,data,user)


@app.post('/api/returns/{return_id}/retry')
def retry_integrated_return(return_id:int,user:User=Depends(identity)):
    return complete_integrated_return(return_id,user)


class CancelSaleIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    reason: str = Field(min_length=3, max_length=160)
    cash_session_id: int | None = Field(default=None, gt=0)
    payment_reference: str | None = Field(default=None, max_length=100)


@app.post('/api/sales/{sale_id}/cancel', status_code=201)
def cancel_sale(sale_id: int, data: CancelSaleIn, idempotency_key: str = Header(min_length=8, max_length=100), user: User = Depends(identity)):
    require(user, 'sale_cancel')
    with Session(engine) as db:
        sale = db.scalar(select(Sale).where(Sale.id == sale_id, Sale.empresa_id == user.empresa_id))
        if sale is None:
            raise HTTPException(404, 'Venta no encontrada')
        branch_for(db, user, sale.branch_id)
        lines = [ReturnLine(sale_item_id=x.id, quantity=x.quantity, restock=True) for x in sale.items]
    return process_return(sale_id, ReturnIn(**data.model_dump(), items=lines), idempotency_key, user, cancellation=True)

@app.get('/api/reports/summary')
def summary(branch_id: int, user: User = Depends(identity)):
    require(user, 'report')
    empresa = user.empresa_id
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        rows = db.scalars(select(Sale).where(Sale.empresa_id == empresa, Sale.branch_id == branch_id)).all()
        returns = db.scalars(select(SaleReturn).join(Sale).where(SaleReturn.status == 'completed', Sale.empresa_id == empresa, Sale.branch_id == branch_id)).all()
        methods = {s.id: s.payment_method for s in rows}
        refunded = sum((r.total for r in returns), Decimal('0'))
        by_method = {method: str(money(sum((s.total for s in rows if s.payment_method == method), Decimal('0')) - sum((r.total for r in returns if methods[r.sale_id] == method), Decimal('0')))) for method in ('cash', 'card', 'transfer', 'mercado_pago')}
        return {'branch_id': branch_id, 'sales_count': len(rows), 'gross_total': str(money(sum((s.total for s in rows), Decimal('0')))), 'returns_total': str(money(refunded)), 'total': str(money(sum((s.total for s in rows), Decimal('0')) - refunded)), 'by_method': by_method}

class SupplierIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    branch_id: int = Field(gt=0)
    name: str = Field(min_length=2, max_length=160)
    phone: str | None = Field(default=None, max_length=30)
    reference: str | None = Field(default=None, max_length=100)

class PurchaseLine(BaseModel):
    product_id: int = Field(gt=0)
    quantity: int = Field(gt=0, le=100000)
    unit_cost: Decimal = Field(ge=0, le=99999999, decimal_places=4)

class PurchaseIn(BaseModel):
    branch_id: int = Field(gt=0)
    supplier_id: int = Field(gt=0)
    reference: str = Field(min_length=1, max_length=100)
    items: list[PurchaseLine] = Field(min_length=1, max_length=100)

class ReceiptIn(BaseModel):
    items: list[Line] = Field(min_length=1, max_length=100)

class CountIn(BaseModel):
    branch_id: int = Field(gt=0)
    product_id: int = Field(gt=0)
    expected: int = Field(ge=0, le=100000000)
    counted: int = Field(ge=0, le=100000000)
    reason: str = Field(min_length=3, max_length=160)

@app.get('/api/suppliers')
def suppliers(branch_id: int, user: User = Depends(identity)):
    require_any(user, 'purchase_write', 'purchase_read')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        return [{'id': s.id, 'name': s.name, 'phone': s.phone, 'reference': s.reference}
                for s in db.scalars(select(Supplier).where(Supplier.empresa_id == user.empresa_id).order_by(Supplier.name))]

@app.post('/api/suppliers', status_code=201)
def create_supplier(data: SupplierIn, user: User = Depends(identity)):
    require(user, 'purchase_write')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        supplier = Supplier(empresa_id=user.empresa_id, **data.model_dump(exclude={'branch_id'}))
        db.add(supplier)
        db.flush()
        db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, actor_id=user.id, action='supplier_created', record_id=supplier.id))
        result = {'id': supplier.id, 'name': supplier.name}
        db.commit()
        return result


@app.put('/api/suppliers/{supplier_id}')
def update_supplier(supplier_id: int, data: SupplierIn, user: User = Depends(identity)):
    require(user, 'purchase_write')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        supplier = db.scalar(select(Supplier).where(Supplier.id == supplier_id, Supplier.empresa_id == user.empresa_id))
        if supplier is None:
            raise HTTPException(404, 'Proveedor no encontrado')
        for key, value in data.model_dump(exclude={'branch_id'}).items():
            setattr(supplier, key, value)
        db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, actor_id=user.id, action='supplier_updated', record_id=supplier.id))
        db.commit()
        return {'id': supplier.id, 'name': supplier.name}

@app.get('/api/purchases/{purchase_id}/receipts')
def purchase_receipts(purchase_id: int, user: User = Depends(identity)):
    import json
    require_any(user, 'purchase_read', 'purchase_write')
    with Session(engine) as db:
        purchase = db.scalar(select(Purchase).where(Purchase.id == purchase_id, Purchase.empresa_id == user.empresa_id))
        if purchase is None:
            raise HTTPException(404, 'Compra no encontrada')
        branch_for(db, user, purchase.branch_id)
        return [{'id': x.id, 'actor_id': x.actor_id, 'created_at': x.created_at.isoformat(),
                'items': [{'product_id': pid, 'quantity': qty} for pid, qty in json.loads(x.payload)]}
                for x in db.scalars(select(PurchaseReceipt).where(PurchaseReceipt.purchase_id == purchase.id).order_by(PurchaseReceipt.id))]


def purchase_view(db, purchase):
    supplier = db.get(Supplier, purchase.supplier_id)
    return {'id': purchase.id, 'branch_id': purchase.branch_id, 'supplier_id': purchase.supplier_id,
            'supplier_name': supplier.name, 'reference': purchase.reference, 'status': purchase.status,
            'created_at': purchase.created_at.isoformat(), 'actor_id': purchase.actor_id,
            'total_cost': str(money(sum((x.unit_cost * x.quantity for x in purchase.items), Decimal('0')))),
            'items': [{'product_id': x.product_id, 'name': x.name, 'quantity': x.quantity,
                       'received': x.received, 'pending': x.quantity - x.received, 'unit_cost': str(x.unit_cost)} for x in purchase.items]}

@app.get('/api/purchases')
def purchases(branch_id: int, user: User = Depends(identity)):
    require_any(user, 'purchase_write', 'purchase_read')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        rows = db.scalars(select(Purchase).where(Purchase.branch_id == branch_id, Purchase.empresa_id == user.empresa_id).order_by(Purchase.id.desc()).limit(100))
        return [purchase_view(db, p) for p in rows]

@app.post('/api/purchases', status_code=201)
def create_purchase(data: PurchaseIn, idempotency_key: str = Header(min_length=8, max_length=100), user: User = Depends(identity)):
    require(user, 'purchase_write')
    ids = [x.product_id for x in data.items]
    if len(ids) != len(set(ids)) or not data.reference.strip():
        raise HTTPException(422, 'Productos duplicados o referencia vacía')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        db.scalar(select(Branch).where(Branch.id == data.branch_id).with_for_update())
        previous = db.scalar(select(Purchase).where(Purchase.empresa_id == user.empresa_id, Purchase.request_key == idempotency_key))
        if previous:
            if (previous.branch_id, previous.supplier_id, previous.reference,
                sorted((x.product_id, x.quantity, x.unit_cost) for x in previous.items)) != (data.branch_id, data.supplier_id,
                data.reference.strip(), sorted((x.product_id, x.quantity, x.unit_cost) for x in data.items)):
                raise HTTPException(409, 'Clave de compra utilizada con otros datos')
            return {**purchase_view(db, previous), 'replayed': True}
        supplier = db.scalar(select(Supplier).where(Supplier.id == data.supplier_id, Supplier.empresa_id == user.empresa_id))
        products = db.scalars(select(Product).where(Product.id.in_(ids), Product.empresa_id == user.empresa_id, Product.active == True)).all()
        if supplier is None or len(products) != len(ids):
            raise HTTPException(404, 'Proveedor o productos fuera de esta empresa, inexistentes o inactivos')
        names = {p.id: p.name for p in products}
        purchase = Purchase(empresa_id=user.empresa_id, branch_id=data.branch_id, supplier_id=supplier.id,
                    reference=data.reference.strip(), request_key=idempotency_key, actor_id=user.id,
                    items=[PurchaseItem(product_id=x.product_id, name=names[x.product_id], quantity=x.quantity, unit_cost=x.unit_cost) for x in data.items])
        db.add(purchase)
        try:
            db.flush()
            db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, action='purchase_ordered', record_id=purchase.id, actor_id=user.id))
            result = {**purchase_view(db, purchase), 'replayed': False}
            db.commit()
            return result
        except IntegrityError:
            raise HTTPException(409, 'Compra duplicada; consulta la lista de compras')

@app.post('/api/purchases/{purchase_id}/receive', status_code=201)
def receive_purchase(purchase_id: int, data: ReceiptIn, idempotency_key: str = Header(min_length=8, max_length=100), user: User = Depends(identity)):
    import json
    require_any(user, 'purchase_write', 'purchase_receive')
    ids = [x.product_id for x in data.items]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, 'Productos duplicados en recepción')
    payload = json.dumps(sorted((x.product_id, x.quantity) for x in data.items), separators=(',', ':'))
    with Session(engine) as db:
        purchase = db.scalar(select(Purchase).where(Purchase.id == purchase_id, Purchase.empresa_id == user.empresa_id))
        if purchase is None:
            raise HTTPException(404, 'Compra no encontrada')
        branch_for(db, user, purchase.branch_id)
        db.scalar(select(Branch).where(Branch.id == purchase.branch_id).with_for_update())
        db.refresh(purchase, with_for_update=True)
        previous = db.scalar(select(PurchaseReceipt).where(PurchaseReceipt.empresa_id == user.empresa_id, PurchaseReceipt.request_key == idempotency_key))
        if previous:
            if previous.purchase_id != purchase_id or previous.payload != payload:
                raise HTTPException(409, 'Clave de recepción utilizada con otros datos')
            return {'id': previous.id, 'purchase_id': purchase.id, 'status': purchase.status, 'replayed': True}
        lines = {x.product_id: x for x in purchase.items}
        for item in data.items:
            if item.product_id not in lines or item.quantity > lines[item.product_id].quantity - lines[item.product_id].received:
                raise HTTPException(409, 'La recepción excede la cantidad pendiente o incluye un producto ajeno')
        stocks = {x.product_id: x for x in db.scalars(select(Stock).where(Stock.product_id.in_(ids), Stock.branch_id == purchase.branch_id).order_by(Stock.product_id).with_for_update())}
        receipt = PurchaseReceipt(empresa_id=user.empresa_id, purchase_id=purchase.id, actor_id=user.id, request_key=idempotency_key, payload=payload)
        db.add(receipt)
        try:
            db.flush()
            for item in data.items:
                line = lines[item.product_id]
                stock = stocks.get(item.product_id)
                if stock is None:
                    stock = Stock(product_id=item.product_id, branch_id=purchase.branch_id, quantity=0, average_cost=Decimal('0'))
                    db.add(stock)
                stock.average_cost = weighted_cost(stock.quantity, stock.average_cost, item.quantity, line.unit_cost)
                stock.quantity += item.quantity
                line.received += item.quantity
                db.add(StockMovement(empresa_id=user.empresa_id, branch_id=purchase.branch_id, product_id=item.product_id,
                        change=item.quantity, reason='purchase_received', reference_id=receipt.id, actor_id=user.id))
            purchase.status = 'received' if all(x.received == x.quantity for x in purchase.items) else 'partial'
            db.add(Audit(empresa_id=user.empresa_id, branch_id=purchase.branch_id, action='purchase_received', record_id=receipt.id, actor_id=user.id))
            result = {'id': receipt.id, 'purchase_id': purchase.id, 'status': purchase.status, 'replayed': False}
            db.commit()
            return result
        except IntegrityError:
            raise HTTPException(409, 'Recepción duplicada; consulta la compra')

@app.get('/api/stock/counts')
def inventory_counts(branch_id: int, user: User = Depends(identity)):
    require_any(user, 'stock_write', 'report')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        return [{'id': x.id, 'product_id': x.product_id, 'expected': x.expected, 'counted': x.counted,
                 'difference': x.counted - x.expected, 'reason': x.reason, 'actor_id': x.actor_id, 'created_at': x.created_at.isoformat()}
                for x in db.scalars(select(InventoryCount).where(InventoryCount.branch_id == branch_id,
                InventoryCount.empresa_id == user.empresa_id).order_by(InventoryCount.id.desc()).limit(100))]

@app.post('/api/stock/counts', status_code=201)
def reconcile_count(data: CountIn, idempotency_key: str = Header(min_length=8, max_length=100), user: User = Depends(identity)):
    require(user, 'count_write')
    if len(data.reason.strip()) < 3:
        raise HTTPException(422, 'Indica el motivo de conciliación')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        db.scalar(select(Branch).where(Branch.id == data.branch_id).with_for_update())
        previous = db.scalar(select(InventoryCount).where(InventoryCount.empresa_id == user.empresa_id, InventoryCount.request_key == idempotency_key))
        if previous:
            if (previous.branch_id, previous.product_id, previous.expected, previous.counted, previous.reason) != (data.branch_id,
                    data.product_id, data.expected, data.counted, data.reason.strip()):
                raise HTTPException(409, 'Clave de conteo utilizada con otros datos')
            return {'id': previous.id, 'difference': previous.counted - previous.expected, 'replayed': True}
        row = db.execute(select(Product, Stock).join(Stock).where(Product.id == data.product_id,
            Product.empresa_id == user.empresa_id, Stock.branch_id == data.branch_id).with_for_update()).first()
        if row is None:
            raise HTTPException(404, 'Producto sin inventario en esta sucursal')
        product, stock = row
        if stock.quantity != data.expected:
            raise HTTPException(409, 'El inventario cambió durante el conteo; actualiza y vuelve a contar')
        count = InventoryCount(empresa_id=user.empresa_id, branch_id=data.branch_id, product_id=product.id, actor_id=user.id,
                               expected=data.expected, counted=data.counted, reason=data.reason.strip(), request_key=idempotency_key)
        if data.counted < reserved_quantity(db, data.branch_id, data.product_id):
            raise HTTPException(409, 'El conteo afecta existencias reservadas para cobros pendientes')
        stock.quantity = data.counted
        db.add(count)
        try:
            db.flush()
            db.add(StockMovement(empresa_id=user.empresa_id, branch_id=data.branch_id, product_id=product.id,
                   actor_id=user.id, change=data.counted - data.expected, reason='count_reconciled', reference_id=count.id))
            db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, actor_id=user.id, action='count_reconciled', record_id=count.id))
            result = {'id': count.id, 'difference': data.counted - data.expected, 'replayed': False}
            db.commit()
            return result
        except IntegrityError:
            raise HTTPException(409, 'Conteo duplicado; consulta el inventario')

class StockMinimumIn(BaseModel):
    branch_id: int = Field(gt=0)
    minimum: int = Field(ge=0, le=1000000)

@app.put('/api/stock/{product_id}/minimum')
def set_minimum(product_id: int, data: StockMinimumIn, user: User = Depends(identity)):
    require(user, 'stock_write')
    with Session(engine) as db:
        branch_for(db, user, data.branch_id)
        stock = db.scalar(select(Stock).join(Product).where(Stock.product_id == product_id,
            Stock.branch_id == data.branch_id, Product.empresa_id == user.empresa_id).with_for_update())
        if stock is None:
            raise HTTPException(404, 'Producto no registrado en esta sucursal')
        stock.minimum = data.minimum
        db.add(Audit(empresa_id=user.empresa_id, branch_id=data.branch_id, actor_id=user.id, action='stock_minimum', record_id=stock.id))
        db.commit()
        return {'product_id': product_id, 'minimum': data.minimum}

@app.get('/api/stock/alerts')
def stock_alerts(branch_id: int, user: User = Depends(identity)):
    require_any(user, 'stock_write', 'report')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        rows = db.execute(select(Product, Stock).join(Stock).where(Product.empresa_id == user.empresa_id,
            Product.active == True, Stock.branch_id == branch_id, Stock.minimum > 0, Stock.quantity <= Stock.minimum).order_by(Stock.quantity, Product.name)).all()
        return [product_view(p, st) for p, st in rows]

class UserUpdateIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    username: str = Field(min_length=3, max_length=80)
    role: str
    active: bool = True
    branch_ids: list[int] = Field(min_length=1)
    password: str | None = Field(default=None, min_length=12, max_length=200)

@app.get('/api/users')
def list_users(user: User = Depends(identity)):
    require(user, 'users_write')
    with Session(engine) as db:
        return [{'id': u.id, 'username': u.username, 'role': u.role, 'active': u.active,
                 'branch_ids': list(db.scalars(select(UserBranch.branch_id).where(UserBranch.user_id == u.id)))}
                for u in db.scalars(select(User).where(User.empresa_id == user.empresa_id).order_by(User.username))]

@app.put('/api/users/{user_id}')
def update_user(user_id: int, data: UserUpdateIn, user: User = Depends(identity)):
    require(user, 'users_write')
    if data.role not in ROLE_ACTIONS:
        raise HTTPException(422, 'Rol inválido')
    with Session(engine) as db:
        # Serialize administrative edits across this company, including last-admin checks.
        db.scalars(select(Branch).where(Branch.empresa_id == user.empresa_id).order_by(Branch.id).with_for_update()).all()
        target = db.scalar(select(User).where(User.id == user_id, User.empresa_id == user.empresa_id).with_for_update())
        if target is None:
            raise HTTPException(404, 'Usuario no encontrado')
        if data.username != target.username and any(c.isspace() for c in data.username):
            raise HTTPException(422, 'El usuario no debe contener espacios')
        ids = set(data.branch_ids)
        valid = db.scalars(select(Branch.id).where(Branch.empresa_id == user.empresa_id, Branch.id.in_(ids))).all()
        if len(valid) != len(data.branch_ids):
            raise HTTPException(422, 'Sucursales inválidas o duplicadas')
        if target.id == user.id and (not data.active or data.role != 'admin_general'):
            raise HTTPException(409, 'No puedes desactivar ni quitar tu propio rol de administración')
        if target.role == 'admin_general' and target.active and (not data.active or data.role != 'admin_general'):
            others = db.scalars(select(User.id).where(User.empresa_id == user.empresa_id, User.role == 'admin_general', User.active == True, User.id != target.id)).all()
            if not others:
                raise HTTPException(409, 'Debe permanecer un administrador general activo')
        if not data.active or data.role != target.role or ids != set(db.scalars(select(UserBranch.branch_id).where(UserBranch.user_id == target.id))):
            if db.scalar(select(CashSession.id).where(CashSession.cashier_id == target.id, CashSession.status == 'open')):
                raise HTTPException(409, 'Cierra los turnos abiertos antes de cambiar rol, sucursales o desactivar')
        assignments = list(db.scalars(select(UserBranch).where(UserBranch.user_id == target.id)))
        target.username, target.role, target.active = data.username, data.role, data.active
        if data.password:
            target.password_hash = password_hash.hash(data.password)
        target.token_version += 1
        for assignment in assignments:
            db.delete(assignment)
        try:
            db.flush()
        except IntegrityError:
            raise HTTPException(409, 'Usuario duplicado')
        db.add_all(UserBranch(user_id=target.id, branch_id=bid) for bid in ids)
        db.add(Audit(empresa_id=user.empresa_id, branch_id=min(ids), actor_id=user.id, action='user_updated', record_id=target.id))
        try:
            db.commit()
        except IntegrityError:
            raise HTTPException(409, 'Usuario duplicado')
        return {'id': target.id, 'username': target.username, 'session_revoked': True}


def build_report(db, user, branch_id, start, end, cashier_id):
    branch_for(db, user, branch_id)
    if end < start or (end - start).days > 365:
        raise HTTPException(422, 'Selecciona un periodo de hasta 366 días en orden válido')
    tz = ZoneInfo('America/Mexico_City')
    lower = datetime.combine(start, datetime.min.time(), tz).astimezone(timezone.utc).replace(tzinfo=None)
    upper = datetime.combine(end + timedelta(days=1), datetime.min.time(), tz).astimezone(timezone.utc).replace(tzinfo=None)
    if cashier_id:
        if not db.scalar(select(User.id).where(User.id == cashier_id, User.empresa_id == user.empresa_id)):
            raise HTTPException(404, 'Cajero no encontrado')
    sale_query = select(Sale).where(Sale.empresa_id == user.empresa_id, Sale.branch_id == branch_id, Sale.created_at >= lower, Sale.created_at < upper)
    return_query = select(SaleReturn).join(Sale).where(SaleReturn.status == 'completed', Sale.empresa_id == user.empresa_id, Sale.branch_id == branch_id, SaleReturn.created_at >= lower, SaleReturn.created_at < upper)
    if cashier_id:
        sale_query = sale_query.where(Sale.actor_id == cashier_id)
        return_query = return_query.where(Sale.actor_id == cashier_id)
    sales = db.scalars(sale_query.order_by(Sale.created_at).limit(50001)).all()
    returns = db.scalars(return_query.order_by(SaleReturn.created_at).limit(50001)).all()
    if len(sales) > 50000 or len(returns) > 50000:
        raise HTTPException(422, 'Demasiados registros; reduce el periodo')
    net, cost, refunded = Decimal('0'), Decimal('0'), Decimal('0')
    products, events = {}, []
    def entry(item):
        return products.setdefault(item.product_id, {'product_id': item.product_id, 'name': item.name, 'units': 0, 'net': Decimal('0')})
    for sale in sales:
        net += sale.subtotal
        for item in sale.items:
            c = item.cost * item.quantity
            cost += c
            p = entry(item); p['units'] += item.quantity
            p['net'] += item.net_amount if item.net_amount is not None else item.unit_price * item.quantity
        events.append({'type': 'Venta', 'id': sale.id, 'folio': sale.folio, 'created_at': sale.created_at.isoformat(), 'cashier_id': sale.actor_id,
                       'method': sale.payment_method, 'amount': str(sale.total)})
    for r in returns:
        sale = db.get(Sale, r.sale_id);refunded += r.total
        for line in r.items:
            item = db.get(SaleItem, line.sale_item_id)
            original = refundable_line_total(sale, item)
            # Allocate the original line net in proportion to the recorded refund.
            original_net = item.net_amount if item.net_amount is not None else item.unit_price * item.quantity
            returned_net = original_net * line.total / original if original else Decimal('0')
            net -= returned_net
            if line.restock:
                cost -= item.cost * line.quantity
            p = entry(item);p['units'] -= line.quantity;p['net'] -= returned_net
        events.append({'type': 'Cancelación' if r.kind == 'cancellation' else 'Devolución', 'id': r.id, 'folio': sale.folio, 'created_at': r.created_at.isoformat(),
                       'cashier_id': sale.actor_id, 'method': sale.payment_method, 'amount': str(-r.total)})
    gross = sum((x.total for x in sales), Decimal('0'))
    return {'branch_id': branch_id, 'start': start.isoformat(), 'end': end.isoformat(), 'cashier_id': cashier_id,
            'timezone': 'America/Mexico_City', 'sales_count': len(sales), 'return_count': len(returns),
            'gross': str(money(gross)), 'refunds': str(money(refunded)), 'total': str(money(gross - refunded)),
            'net_before_tax': str(money(net)), 'cost': str(money(cost)), 'profit': str(money(net - cost)),
            'missing_cost_lines': sum(x.cost == 0 for sale in sales for x in sale.items),
            'products': [{**p, 'net': str(money(p['net']))} for p in sorted(products.values(), key=lambda p: (-p['units'], p['name']))],
            'events': sorted(events, key=lambda e: e['created_at'])}

@app.get('/api/reports/cashiers')
def report_cashiers(branch_id: int, user: User = Depends(identity)):
    require(user, 'report')
    with Session(engine) as db:
        branch_for(db, user, branch_id)
        ids = db.scalars(select(Sale.actor_id).where(Sale.empresa_id == user.empresa_id, Sale.branch_id == branch_id, Sale.actor_id.is_not(None)).distinct()).all()
        return [{'id': u.id, 'name': u.username} for u in db.scalars(select(User).where(User.empresa_id == user.empresa_id, User.id.in_(ids)).order_by(User.username))]

@app.get('/api/reports/sales')
def sales_report(branch_id: int, start: date, end: date, cashier_id: int | None = None, user: User = Depends(identity)):
    require(user, 'report')
    with Session(engine) as db:
        return build_report(db, user, branch_id, start, end, cashier_id)

@app.get('/api/reports/sales.xlsx')
def sales_report_excel(branch_id: int, start: date, end: date, cashier_id: int | None = None, user: User = Depends(identity)):
    require(user, 'report')
    with Session(engine) as db:
        report = build_report(db, user, branch_id, start, end, cashier_id)
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = 'Resumen'
    ws.append(['LI Punto de Venta', 'Reporte por fecha de operación'])
    for label, value in [('Sucursal', branch_id), ('Desde', report['start']), ('Hasta', report['end']), ('Zona horaria', report['timezone']),
                         ('Ventas brutas', report['gross']), ('Reembolsos', report['refunds']), ('Total neto', report['total']), ('Utilidad estimada', report['profit']), ('Partidas sin costo', report['missing_cost_lines'])]:
        ws.append([label, float(Decimal(value)) if label in ('Ventas brutas', 'Reembolsos', 'Total neto', 'Utilidad estimada') else value])
    for title, headers, rows in [('Movimientos', ['Tipo','ID','Folio','Fecha UTC','Cajero','Pago','Importe'], [[float(Decimal(e[k])) if k == 'amount' else e[k] for k in ('type','id','folio','created_at','cashier_id','method','amount')] for e in report['events']]),
                                 ('Productos', ['ID','Producto','Unidades netas','Venta neta sin impuesto'], [[float(Decimal(p[k])) if k == 'net' else p[k] for k in ('product_id','name','units','net')] for p in report['products']])]:
        sheet = wb.create_sheet(title);sheet.append(headers)
        for row in rows:
            sheet.append(row)
            # Prevent formula injection from names and other stored strings.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):cell.data_type = 's'
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    from openpyxl.styles import Font, PatternFill
    for sheet in wb:
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='252B36')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=min(55,max(16,max(len(str(c.value or '')) for c in column)+2))
    output=BytesIO();wb.save(output)
    return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition':'attachment; filename="LI_Reporte_Ventas.xlsx"'})

from . import mercado_pago as mp


def intent_view(intent):
    return {'id': intent.id, 'amount': str(intent.amount), 'status': intent.status, 'actor_id': intent.actor_id,
            'checkout_url': None if intent.cancel_requested_at else intent.checkout_url, 'payment_id': intent.payment_id,
            'cancellation_pending': bool(intent.cancel_requested_at and not intent.cancelled_at),
            'cancelled_at': intent.cancelled_at.isoformat() if intent.cancelled_at else None,
            'original_cash_session_id': json.loads(intent.payload).get('cash_session_id'),
            'delivery_cash_session_id': intent.delivery_cash_session_id,
            'delivered_at': intent.delivered_at.isoformat() if intent.delivered_at else None,
            'created_at': intent.created_at.isoformat(), 'reserved': intent.reservation_active,
            'review_reason': intent.review_reason, 'resolved_by':intent.resolved_by,
            'resolved_at':intent.resolved_at.isoformat() if intent.resolved_at else None}


@app.post('/api/payments/checkout', status_code=201)
def create_payment_checkout(data: SaleIn, idempotency_key: str = Header(min_length=8,max_length=100), user: User = Depends(identity)):
    require(user,'sale');config=mp.settings()
    if user.empresa_id != config['company']:
        raise HTTPException(403,'Cuenta de Mercado Pago no configurada para esta empresa')
    if data.discount_percent:
        raise HTTPException(422,'El checkout integrado inicial no admite descuentos; usa la venta manual autorizada')
    quoted=sale_quote(data,user)
    normalized=data.model_copy(update={'payment_method':'mercado_pago','paid':Decimal(quoted['total']),'payment_intent_id':None,'approval':None})
    payload=json.dumps(normalized.model_dump(mode='json',exclude={'approval','payment_intent_id'}),sort_keys=True)
    with Session(engine) as db:
        branch_for(db,user,data.branch_id)
        db.scalar(select(Branch).where(Branch.id==data.branch_id).with_for_update())
        cash=db.scalar(select(CashSession).where(CashSession.id==data.cash_session_id,CashSession.branch_id==data.branch_id,CashSession.status=='open'))
        if not cash:raise HTTPException(409,'Abre tu turno antes de generar un cobro')
        cash_access(db,user,cash,owner=True)
        intent=db.scalar(select(PaymentIntent).where(PaymentIntent.empresa_id==user.empresa_id,PaymentIntent.request_key==idempotency_key))
        if intent:
            if intent.actor_id!=user.id or intent.payload!=payload:raise HTTPException(409,'Clave utilizada para otro cobro')
            if intent.checkout_url:return intent_view(intent)
            raise HTTPException(409,'Creación previa pendiente de conciliación; no generes otro cobro')
        snapshot = {}
        snapshot_total = Decimal(0)
        for line in data.items:
            product = db.scalar(select(Product).where(Product.id == line.product_id, Product.empresa_id == user.empresa_id))
            if product is None or not product.active:
                raise HTTPException(409, "Producto inválido para checkout")
            snapshot[str(product.id)] = {k: str(getattr(product,k)) if k in ("price","tax_rate") else getattr(product,k)
                for k in ("id","name","price","unit","tax_rate","tax_exempt","price_includes_tax")}
            net, tax, _ = price_line(product, line.quantity, Decimal(0))
            snapshot_total += net + tax
            stock=db.scalar(select(Stock).where(Stock.product_id==line.product_id,Stock.branch_id==data.branch_id))
            if not stock or stock.quantity-reserved_quantity(db,data.branch_id,line.product_id)<line.quantity:raise HTTPException(409,'Existencias insuficientes')
        if snapshot_total != Decimal(quoted["total"]):
            raise HTTPException(409, "El precio cambió al crear el checkout; vuelve a cotizar")
        intent=PaymentIntent(id=secrets.token_urlsafe(24),empresa_id=user.empresa_id,branch_id=data.branch_id,actor_id=user.id,
            payload=payload,amount=Decimal(quoted['total']),request_key=idempotency_key,reservation_active=True,price_snapshot=json.dumps(snapshot))
        db.add(intent);db.commit();intent_id=intent.id
    result=mp.call('POST','/checkout/preferences',{'external_reference':intent_id,
        'items':[{'id':intent_id,'title':'LI Punto de Venta · compra','quantity':1,'currency_id':'MXN','unit_price':float(Decimal(quoted['total']))}],
        'notification_url':config['url']+'/api/payments/webhook',
        'back_urls':{s:config['url']+'/' for s in ('success','pending','failure')},'auto_return':'approved',
        'expires':True,'expiration_date_from':datetime.now(timezone.utc).isoformat(),
        'expiration_date_to':(datetime.now(timezone.utc)+timedelta(minutes=30)).isoformat()},key=intent_id)
    url=result.get('init_point' if config['live'] else 'sandbox_init_point')
    from urllib.parse import urlparse
    parsed=urlparse(url or '')
    if parsed.scheme!='https' or not any((parsed.hostname or '').endswith(suffix) for suffix in ('.mercadopago.com','.mercadopago.com.mx')):
        raise HTTPException(502,'URL de checkout inválida; consulta la operación pendiente')
    with Session(engine) as db:
        intent=db.get(PaymentIntent,intent_id);intent.preference_id=str(result['id']);intent.checkout_url=url
        if intent.status=='creating':intent.status='pending'
        db.commit()
        return intent_view(intent)


def apply_provider_payment(payment):
    config=mp.settings()
    with Session(engine) as db:
        intent=db.get(PaymentIntent,str(payment.get('external_reference','')))
        if intent is None:return {'ignored':True}
        if (intent.empresa_id!=config['company'] or str(payment.get('collector_id'))!=config['collector']
            or payment.get('currency_id')!='MXN' or Decimal(str(payment.get('transaction_amount',0)))!=intent.amount
            or payment.get('live_mode') is not config['live']):
            raise HTTPException(409,'Pago no coincide con empresa, importe, moneda o entorno')
        db.scalar(select(Branch).where(Branch.id==intent.branch_id).with_for_update())
        intent=db.scalar(select(PaymentIntent).where(PaymentIntent.id==intent.id).with_for_update().execution_options(populate_existing=True))
        pid=str(payment['id'])
        record_payment_observation(db,intent,payment)
        if intent.cancel_requested_at:
            if not intent.payment_id:
                intent.payment_id = pid
            if not fully_refunded(payment) and payment.get('status') not in ('cancelled', 'rejected') and (intent.cancelled_at or payment.get('status') not in ('pending','in_process','authorized')):
                intent.review_reason = 'Pago durante o después de cancelación: '+pid[:100]+' · '+str(payment.get('status'))[:40]
            db.commit()
            return intent_view(intent)
        if intent.payment_id and intent.payment_id!=pid:
            if fully_refunded(payment) or payment.get('status') in ('cancelled','rejected'):
                db.commit()
                return intent_view(intent)
            intent.review_reason='Otro pago detectado: '+pid[:100]
            db.commit()
            raise HTTPException(409,'Hay otro pago para este checkout; requiere conciliación')
        if intent.status=='completed':
            sale=db.scalar(select(Sale).where(Sale.empresa_id==intent.empresa_id,Sale.request_key=='mp-sale-'+intent.id))
            local_refunded=sum((r.total for r in db.scalars(select(SaleReturn).where(SaleReturn.sale_id==sale.id,SaleReturn.status=='completed'))),Decimal(0)) if sale else Decimal(0)
            if payment.get('status') not in ('approved','refunded') or Decimal(str(payment.get('transaction_amount_refunded',0)))!=local_refunded:
                intent.review_reason='Pago entregado: estado '+str(payment.get('status'))+'; reembolso '+str(payment.get('transaction_amount_refunded',0))
                db.commit()
                return intent_view(intent)
            db.commit()
            return intent_view(intent)
        intent.payment_id=pid;intent.status='refunded' if Decimal(str(payment.get('transaction_amount_refunded',0)))>0 else str(payment.get('status','pending'))[:40]
        if intent.status in ('refunded', 'charged_back'):
            intent.review_reason='Requiere conciliación financiera: '+intent.status
        try:db.commit()
        except IntegrityError:raise HTTPException(409,'Pago ya asociado a otra operación')
        return intent_view(intent)


@app.post('/api/payments/webhook')
async def payment_webhook(request: Request):
    data_id=request.query_params.get('data.id','')
    mp.verify_signature(request.headers.get('x-signature'),request.headers.get('x-request-id'),data_id)
    if not data_id.isdigit():raise HTTPException(422,'Identificador de pago inválido')
    body=await request.json()
    if body.get('type')!='payment':return {'ignored':True}
    # Neither redirect parameters nor the notification body confirm payment.
    from starlette.concurrency import run_in_threadpool
    payment=await run_in_threadpool(mp.call,'GET','/v1/payments/'+data_id)
    return await run_in_threadpool(apply_provider_payment,payment)


@app.get('/api/payments')
def list_payments(branch_id:int,user:User=Depends(identity)):
    require_any(user,'sale','report')
    with Session(engine) as db:
        branch_for(db,user,branch_id)
        query=select(PaymentIntent).where(PaymentIntent.empresa_id==user.empresa_id,PaymentIntent.branch_id==branch_id)
        if user.role not in ('admin_general','admin_sucursal'):query=query.where(PaymentIntent.actor_id==user.id)
        return [intent_view(i) for i in db.scalars(query.order_by(PaymentIntent.created_at.desc()).limit(100))]


@app.post('/api/payments/{intent_id}/reconcile')
def reconcile_integrated_payment(intent_id: str, user: User = Depends(identity)):
    require_any(user, 'sale', 'report')
    from urllib.parse import urlencode
    config = mp.settings()
    with Session(engine) as db:
        intent = db.scalar(select(PaymentIntent).where(PaymentIntent.id == intent_id,
            PaymentIntent.empresa_id == user.empresa_id))
        if intent is None or (user.role not in ('admin_general', 'admin_sucursal') and intent.actor_id != user.id):
            raise HTTPException(404, 'Cobro no encontrado')
        branch_for(db, user, intent.branch_id)
        if intent.empresa_id != config['company']:
            raise HTTPException(403, 'Cuenta de pagos fuera de esta empresa')
    result = mp.call('GET', '/v1/payments/search?' + urlencode({'external_reference': intent_id, 'limit': 100}))
    payments = result.get('results', [])
    if result.get('paging', {}).get('total', len(payments)) > len(payments):
        raise HTTPException(409, 'Demasiados pagos para conciliación automática')
    for payment in payments:
        if str(payment.get('external_reference')) != intent_id:
            raise HTTPException(502, 'Respuesta del proveedor fuera de esta operación')
        # Consult the authoritative resource rather than trusting search details.
        apply_provider_payment(mp.call('GET', '/v1/payments/' + str(payment['id'])))
    with Session(engine) as db:
        intent = db.get(PaymentIntent, intent_id)
        db.add(Audit(empresa_id=user.empresa_id, branch_id=intent.branch_id,
            actor_id=user.id, action='payment_reconciled', record_id=user.id))
        db.commit()
        return intent_view(intent)


class CancelCheckoutIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    reason: str = Field(min_length=3, max_length=300)


def validate_cancel_payment(payment, intent, config):
    if (str(payment.get('external_reference')) != intent.id
        or str(payment.get('collector_id')) != config['collector']
        or payment.get('currency_id') != 'MXN'
        or Decimal(str(payment.get('transaction_amount', 0))) != intent.amount
        or payment.get('live_mode') is not config['live']):
        raise HTTPException(409, 'Pago fuera de la operación; reserva conservada')


def verified_payment_ids(intent_id, known_id):
    from urllib.parse import urlencode
    result = mp.call('GET', '/v1/payments/search?' + urlencode({'external_reference': intent_id, 'limit': 100}))
    rows = result.get('results')
    if not isinstance(rows, list) or result.get('paging', {}).get('total', len(rows)) > len(rows):
        raise HTTPException(409, 'Búsqueda incompleta; reserva conservada')
    ids = set()
    for row in rows:
        if str(row.get('external_reference')) != intent_id or not str(row.get('id', '')).isdigit():
            raise HTTPException(502, 'Respuesta de pagos inválida; reserva conservada')
        ids.add(str(row['id']))
    if known_id:
        ids.add(known_id)
    with Session(engine) as db:
        ids.update(db.scalars(select(PaymentObservation.payment_id).where(PaymentObservation.intent_id==intent_id)))
        ids.update(db.scalars(select(PaymentRefund.payment_id).where(PaymentRefund.intent_id==intent_id)))
    return ids


@app.post('/api/payments/{intent_id}/cancel')
def cancel_payment_checkout(intent_id: str, data: CancelCheckoutIn, user: User = Depends(identity)):
    require(user, 'sale')
    config = mp.settings()
    with Session(engine) as db:
        intent = db.scalar(select(PaymentIntent).where(PaymentIntent.id == intent_id,
            PaymentIntent.empresa_id == user.empresa_id))
        if intent is None or (intent.actor_id != user.id and user.role not in ('admin_general','admin_sucursal')):
            raise HTTPException(404, 'Cobro no encontrado')
        branch_for(db, user, intent.branch_id)
        if intent.empresa_id != config['company']:
            raise HTTPException(403, 'Cuenta de pagos fuera de esta empresa')
        db.scalar(select(Branch).where(Branch.id == intent.branch_id).with_for_update())
        intent = db.scalar(select(PaymentIntent).where(PaymentIntent.id == intent_id)
            .with_for_update().execution_options(populate_existing=True))
        if intent.cancelled_at:
            return intent_view(intent)
        if intent.status in ('approved', 'completed') or intent.review_reason:
            raise HTTPException(409, 'Pago aprobado o con incidencia: requiere conciliación, no cancelación de checkout')
        if not intent.cancel_requested_at:
            intent.cancel_requested_at = datetime.now(timezone.utc)
            intent.cancel_reason = data.reason
            intent.cancel_actor_id = user.id
            intent.status = 'cancel_pending'
            db.add(Audit(empresa_id=user.empresa_id, branch_id=intent.branch_id,
                actor_id=user.id, action='checkout_cancel_requested', record_id=user.id))
        db.commit()
        # Durable cancellation blocks ticket issuance before any remote mutation.
        db.refresh(intent)
        db.expunge(intent)
    preference_id = intent.preference_id
    if not preference_id:
        from urllib.parse import urlencode
        result = mp.call('GET', '/checkout/preferences/search?' + urlencode({'external_reference': intent.id}))
        elements = result.get('elements', [])
        if len(elements) != 1 or result.get('total', len(elements)) != 1:
            raise HTTPException(409, 'Preferencia no localizada de forma única; reserva conservada. Reintenta consultar después')
        preference_id = str(elements[0]['id'])
    from urllib.parse import quote
    path = '/checkout/preferences/' + quote(preference_id, safe='')
    preference = mp.call('GET', path)
    if str(preference.get('external_reference')) != intent.id or str(preference.get('collector_id')) != config['collector']:
        raise HTTPException(409, 'Preferencia fuera de la operación; reserva conservada')
    deadline = datetime.now(timezone.utc) - timedelta(seconds=1)
    mp.call('PUT', path, {'expires': True, 'expiration_date_from': (deadline-timedelta(days=1)).isoformat(),
        'expiration_date_to': deadline.isoformat()}, key='cancel-pref-'+intent.id)
    verified = mp.call('GET', path)
    try:
        expires_at = datetime.fromisoformat(verified.get('expiration_date_to','').replace('Z','+00:00'))
        disabled = (verified.get('expires') is True and expires_at.tzinfo is not None
            and expires_at <= datetime.now(timezone.utc)
            and str(verified.get('external_reference')) == intent.id
            and str(verified.get('collector_id')) == config['collector'])
    except (ValueError, TypeError):
        disabled = False
    if not disabled:
        raise HTTPException(409, 'Proveedor no confirmó cierre del enlace; reserva conservada')
    checked = set()
    # Search again after pending payments are cancelled to catch another attempt.
    for _ in range(2):
        ids = verified_payment_ids(intent.id, intent.payment_id)
        for pid in sorted(ids):
            payment = mp.call('GET', '/v1/payments/'+pid)
            validate_cancel_payment(payment, intent, config)
            if payment.get('status') in ('pending','in_process','authorized'):
                mp.call('PUT', '/v1/payments/'+pid, {'status':'cancelled'}, key='cancel-pay-'+intent.id+'-'+pid)
                payment = mp.call('GET', '/v1/payments/'+pid)
                validate_cancel_payment(payment, intent, config)
            if payment.get('status') not in ('cancelled','rejected') and not fully_refunded(payment):
                apply_provider_payment(payment)
                raise HTTPException(409, 'El pago no está cancelado; reserva conservada y requiere conciliación')
            checked.add(pid)
    with Session(engine) as db:
        db.scalar(select(Branch).where(Branch.id == intent.branch_id).with_for_update())
        current = db.scalar(select(PaymentIntent).where(PaymentIntent.id == intent.id).with_for_update())
        if current.cancelled_at:
            return intent_view(current)
        if current.review_reason or current.status in ('approved','completed') or (current.payment_id and current.payment_id not in checked):
            raise HTTPException(409, 'Cambió el pago durante cancelación; reserva conservada')
        current.preference_id = preference_id
        current.status = 'cancelled'
        current.cancelled_at = datetime.now(timezone.utc)
        current.reservation_active = False
        current.checkout_url = None
        db.add(Audit(empresa_id=user.empresa_id, branch_id=current.branch_id,
            actor_id=user.id, action='checkout_cancelled', record_id=user.id))
        db.commit()
        return intent_view(current)


def scoped_payment_intent(db, user, intent_id):
    intent = db.scalar(select(PaymentIntent).where(PaymentIntent.id == intent_id,
        PaymentIntent.empresa_id == user.empresa_id))
    if intent is None or (intent.actor_id != user.id and user.role not in ('admin_general','admin_sucursal')):
        raise HTTPException(404, 'Cobro no encontrado')
    branch_for(db, user, intent.branch_id)
    if intent.empresa_id != mp.settings()['company']:
        raise HTTPException(403, 'Cuenta de pagos fuera de esta empresa')
    return intent


def refund_view(record):
    return {'id':record.id,'payment_id':record.payment_id,'amount':str(record.amount),
        'status':record.status,'provider_refund_id':record.provider_refund_id,
        'reason':record.reason,'actor_id':record.actor_id,
        'confirmed_at':record.confirmed_at.isoformat() if record.confirmed_at else None}


@app.get('/api/payments/{intent_id}/details')
def integrated_payment_details(intent_id: str, user: User = Depends(identity)):
    require_any(user,'sale','report')
    config=mp.settings()
    with Session(engine) as db:
        intent=scoped_payment_intent(db,user,intent_id)
        db.expunge(intent)
    payments=[]
    for pid in sorted(verified_payment_ids(intent.id,intent.payment_id)):
        payment=mp.call('GET','/v1/payments/'+pid)
        validate_cancel_payment(payment,intent,config)
        payments.append({'id':pid,'status':payment.get('status'),
            'amount':str(intent.amount),'refunded':str(payment.get('transaction_amount_refunded',0)),
            'refundable':bool(user.role in ('admin_general','admin_sucursal')
                and payment.get('status')=='approved' and not fully_refunded(payment)
                and not (intent.status=='completed' and pid==intent.payment_id))})
    with Session(engine) as db:
        return {'intent':intent_view(db.get(PaymentIntent,intent_id)), 'payments':payments,
            'refunds':[refund_view(r) for r in db.scalars(select(PaymentRefund).where(PaymentRefund.intent_id==intent_id))]}


class RefundIncidentIn(BaseModel):
    model_config=ConfigDict(str_strip_whitespace=True,extra='forbid')
    payment_id: str = Field(pattern=r'^[0-9]{1,100}$')
    reason: str = Field(min_length=3,max_length=300)


@app.post('/api/payments/{intent_id}/refund')
def refund_checkout_incident(intent_id: str, data: RefundIncidentIn, user: User = Depends(identity)):
    require(user,'sale_return');config=mp.settings()
    with Session(engine) as db:
        intent=scoped_payment_intent(db,user,intent_id)
        db.expunge(intent)
    payment=mp.call('GET','/v1/payments/'+data.payment_id)
    validate_cancel_payment(payment,intent,config)
    if payment.get('status') not in ('approved','refunded'):
        raise HTTPException(409,'Estado no reembolsable; consulta el proveedor')
    remaining=money(intent.amount-Decimal(str(payment.get('transaction_amount_refunded',0))))
    if remaining<0:
        raise HTTPException(409,'Importe reembolsado inválido')
    with Session(engine) as db:
        db.scalar(select(Branch).where(Branch.id==intent.branch_id).with_for_update())
        current=db.scalar(select(PaymentIntent).where(PaymentIntent.id==intent.id).with_for_update())
        delivered=db.scalar(select(Sale).where(Sale.empresa_id==user.empresa_id,Sale.request_key=='mp-sale-'+intent.id))
        if delivered and (current.payment_id==data.payment_id or not current.payment_id):
            raise HTTPException(409,'Pago de ticket entregado: utiliza devoluciones para conservar inventario y contabilidad')
        operation=db.scalar(select(PaymentRefund).where(PaymentRefund.payment_id==data.payment_id))
        if operation and (operation.empresa_id!=user.empresa_id or operation.intent_id!=intent.id):
            raise HTTPException(409,'Pago asociado a otra operación de reembolso')
        if operation and operation.status=='confirmed' and fully_refunded(payment):
            return refund_view(operation)
        if operation is None:
            import uuid
            operation=PaymentRefund(empresa_id=user.empresa_id,intent_id=intent.id,payment_id=data.payment_id,
                amount=remaining,provider_key=str(uuid.uuid4()),reason=data.reason,actor_id=user.id)
            db.add(operation);db.flush()
            db.add(Audit(empresa_id=user.empresa_id,branch_id=intent.branch_id,actor_id=user.id,
                action='payment_refund_requested',record_id=operation.id))
        record_payment_observation(db,current,payment)
        if not delivered and (not current.payment_id or current.payment_id==data.payment_id) and not current.cancel_requested_at:
            current.cancel_requested_at=datetime.now(timezone.utc)
            current.cancel_reason='Reembolso: '+data.reason[:280]
            current.cancel_actor_id=user.id
            current.status='cancel_pending'
        current.review_reason='Reembolso en conciliación del pago '+data.payment_id
        db.commit();db.refresh(operation);db.expunge(operation)
    try:
        if not fully_refunded(payment):
            if operation.status=='confirmed':
                raise HTTPException(409,'El estado de un reembolso confirmado cambió; requiere revisión')
            if operation.amount<=0:
                raise HTTPException(409,'Importe del reembolso requiere revisión')
            result=mp.call('POST','/v1/payments/'+data.payment_id+'/refunds',
                {'amount':float(operation.amount)},key=operation.provider_key)
            if str(result.get('payment_id'))!=data.payment_id or Decimal(str(result.get('amount',0)))!=operation.amount:
                raise HTTPException(502,'Respuesta de reembolso no coincide; operación conservada para consulta')
            with Session(engine) as db:
                stored=db.get(PaymentRefund,operation.id)
                stored.provider_refund_id=str(result.get('id',''))[:100] or None
                db.commit()
        verified=mp.call('GET','/v1/payments/'+data.payment_id)
        validate_cancel_payment(verified,intent,config)
        with Session(engine) as db:
            db.scalar(select(Branch).where(Branch.id==intent.branch_id).with_for_update())
            record_payment_observation(db,db.get(PaymentIntent,intent.id),verified)
            stored=db.get(PaymentRefund,operation.id)
            if fully_refunded(verified):
                if stored.status!='confirmed':
                    stored.status='confirmed';stored.confirmed_at=datetime.now(timezone.utc)
                    db.add(Audit(empresa_id=user.empresa_id,branch_id=intent.branch_id,actor_id=user.id,
                        action='payment_refund_confirmed',record_id=stored.id))
            else:
                stored.status='uncertain'
            db.commit()
            return refund_view(stored)
    except HTTPException:
        with Session(engine) as db:
            stored=db.get(PaymentRefund,operation.id)
            if stored.status!='confirmed':stored.status='uncertain'
            db.commit()
        raise


@app.post('/api/payments/{intent_id}/resolve')
def resolve_payment_incident(intent_id: str, user: User = Depends(identity)):
    require(user,'sale_return');config=mp.settings()
    with Session(engine) as db:
        intent=scoped_payment_intent(db,user,intent_id)
        baseline={r.payment_id:r.updated_at for r in db.scalars(select(PaymentObservation).where(PaymentObservation.intent_id==intent.id))}
        db.expunge(intent)
    observations={}
    for pid in sorted(verified_payment_ids(intent.id,intent.payment_id)):
        pay=mp.call('GET','/v1/payments/'+pid)
        validate_cancel_payment(pay,intent,config)
        observations[pid]=pay
    with Session(engine) as db:
        db.scalar(select(Branch).where(Branch.id==intent.branch_id).with_for_update())
        current=db.scalar(select(PaymentIntent).where(PaymentIntent.id==intent.id).with_for_update())
        records=db.scalars(select(PaymentObservation).where(PaymentObservation.intent_id==intent.id)).all()
        if any(baseline.get(r.payment_id)!=r.updated_at for r in records):
            raise HTTPException(409,'Llegó una notificación durante conciliación; consulta de nuevo')
        known={r.payment_id for r in records}
        if not known <= observations.keys() or (current.payment_id and current.payment_id not in observations):
            raise HTTPException(409,'El pago cambió durante la consulta; vuelve a conciliar')
        refunds=db.scalars(select(PaymentRefund).where(PaymentRefund.intent_id==intent.id)).all()
        for refund in refunds:
            pay=observations.get(refund.payment_id)
            if not pay or not fully_refunded(pay):
                raise HTTPException(409,'Hay un reembolso sin confirmar; conserva la incidencia')
        delivered=db.scalar(select(Sale).where(Sale.empresa_id==user.empresa_id,Sale.request_key=='mp-sale-'+intent.id))
        if delivered and not current.payment_id:
            raise HTTPException(409,'Ticket sin pago principal identificado; requiere revisión')
        candidates=[]
        for pid,pay in observations.items():
            if delivered and pid==current.payment_id:
                local_refunded=sum((r.total for r in db.scalars(select(SaleReturn).where(SaleReturn.sale_id==delivered.id,SaleReturn.status=='completed'))),Decimal(0))
                if pay.get('status') not in ('approved','refunded') or Decimal(str(pay.get('transaction_amount_refunded',0)))!=local_refunded:
                    raise HTTPException(409,'El reembolso o contracargo del ticket no coincide con sus devoluciones')
                if db.scalar(select(SaleReturn.id).where(SaleReturn.sale_id==delivered.id,SaleReturn.status!='completed')):
                    raise HTTPException(409,'Hay una devolución pendiente de confirmar')
            elif pay.get('status') not in ('cancelled','rejected') and not fully_refunded(pay):
                if (not delivered and not current.cancel_requested_at and pay.get('status')=='approved'
                    and Decimal(str(pay.get('transaction_amount_refunded',0)))==0):
                    candidates.append(pid)
                else:
                    raise HTTPException(409,'Persisten pagos pendientes, aprobados o contracargos; incidencia conservada')
        if len(candidates)>1:
            raise HTTPException(409,'Hay más de un pago aprobado; reembolsa el excedente')
        if candidates:
            current.payment_id=candidates[0];current.status='approved'
        elif not delivered and not current.cancel_requested_at:
            if not observations:
                raise HTTPException(409,'No hay pagos observados: cancela el checkout para cerrar la operación')
            current.cancel_requested_at=datetime.now(timezone.utc)
            current.cancel_actor_id=user.id;current.cancel_reason='Conciliación de pagos sin entrega'
            current.status='cancel_pending'
        for pay in observations.values():
            record_payment_observation(db,current,pay)
        for refund in refunds:
            if refund.status!='confirmed':
                refund.status='confirmed';refund.confirmed_at=datetime.now(timezone.utc)
                db.add(Audit(empresa_id=user.empresa_id,branch_id=intent.branch_id,actor_id=user.id,
                    action='payment_refund_confirmed',record_id=refund.id))
        current.review_reason=None
        current.resolved_by=user.id;current.resolved_at=datetime.now(timezone.utc)
        db.add(Audit(empresa_id=user.empresa_id,branch_id=intent.branch_id,actor_id=user.id,
            action='payment_incident_resolved',record_id=user.id))
        db.commit()
        return intent_view(current)


class ConfirmPaymentIn(BaseModel):
    model_config = ConfigDict(extra='forbid')
    cash_session_id: int | None = Field(default=None, gt=0)


@app.post('/api/payments/{intent_id}/confirm')
def confirm_integrated_sale(intent_id: str, data: ConfirmPaymentIn, user: User = Depends(identity)):
    require(user,'sale')
    with Session(engine) as db:
        intent=db.scalar(select(PaymentIntent).where(PaymentIntent.id==intent_id,PaymentIntent.empresa_id==user.empresa_id,PaymentIntent.actor_id==user.id))
        if intent is None:raise HTTPException(404,'Cobro no encontrado para este cajero')
        branch_for(db,user,intent.branch_id)
        payment_id=intent.payment_id;payload=json.loads(intent.payload)
        delivered = db.scalar(select(Sale).where(Sale.empresa_id == user.empresa_id,
            Sale.request_key == 'mp-sale-'+intent_id))
        if delivered:
            # A replay always uses the recorded delivery shift, even after closing it.
            payload['cash_session_id'] = delivered.cash_session_id
        elif data.cash_session_id is not None:
            payload['cash_session_id'] = data.cash_session_id
        if not payment_id:raise HTTPException(409,'Esperando notificación de Mercado Pago')
    updated=apply_provider_payment(mp.call('GET','/v1/payments/'+payment_id))
    if updated.get('review_reason') or updated['status'] not in ('approved','completed'):raise HTTPException(409,'El proveedor no confirma pago aprobado')
    return sell(SaleIn(**payload,payment_intent_id=intent_id),'mp-sale-'+intent_id,user)

if os.getenv('APP_ENV', 'demo') != 'production':
    if 'audit_logs' in inspect(engine).get_table_names() and 'actor_id' not in {c['name'] for c in inspect(engine).get_columns('audit_logs')}:
        raise RuntimeError('Base demo anterior: ejecuta python scripts/upgrade_demo_sqlite.py pos.db')
    if 'products' in inspect(engine).get_table_names() and 'barcode' not in {c['name'] for c in inspect(engine).get_columns('products')}:
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    if 'products' in inspect(engine).get_table_names() and ('sale_returns' not in inspect(engine).get_table_names() or 'kind' not in {c['name'] for c in inspect(engine).get_columns('sale_returns')}):
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    if 'stock' in inspect(engine).get_table_names() and 'minimum' not in {c['name'] for c in inspect(engine).get_columns('stock')}:
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    if 'stock' in inspect(engine).get_table_names() and 'payment_intents' not in inspect(engine).get_table_names():
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    if 'payment_intents' in inspect(engine).get_table_names() and 'reservation_active' not in {c['name'] for c in inspect(engine).get_columns('payment_intents')}:
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    if 'payment_intents' in inspect(engine).get_table_names() and 'cancel_requested_at' not in {c['name'] for c in inspect(engine).get_columns('payment_intents')}:
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    if 'payment_intents' in inspect(engine).get_table_names() and 'delivery_cash_session_id' not in {c['name'] for c in inspect(engine).get_columns('payment_intents')}:
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    if 'payment_intents' in inspect(engine).get_table_names() and not {'payment_refunds','payment_observations'} <= set(inspect(engine).get_table_names()):
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    if 'sale_returns' in inspect(engine).get_table_names() and 'status' not in {c['name'] for c in inspect(engine).get_columns('sale_returns')}:
        raise RuntimeError('Base anterior: ejecuta python scripts/migrate_local.py antes de iniciar')
    Base.metadata.create_all(engine)  # Only for a fresh local demo DB.
    with Session(engine) as db:
        if not db.scalar(select(Branch.id).where(Branch.empresa_id == 1).limit(1)):
            db.add_all(Branch(empresa_id=1, name=name) for name in ('Zamora', 'Zacapu', 'Uruapan', '20 de Noviembre', 'Maravatío', 'CDMX'))
            db.commit()
        for branch in db.scalars(select(Branch)):
            if not db.scalar(select(CashRegister.id).where(CashRegister.branch_id == branch.id)):
                db.add(CashRegister(empresa_id=branch.empresa_id, branch_id=branch.id, name='Caja 1'))
        db.commit()
frontend = Path(__file__).resolve().parents[2] / 'frontend' / 'public'
app.mount('/assets', StaticFiles(directory=frontend), name='assets')
@app.get('/')
def home():
    return FileResponse(frontend / 'index.html')


