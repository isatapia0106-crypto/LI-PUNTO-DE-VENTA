"""Mínimos, usuarios y reportes de fecha de operación."""
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
import uuid
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_sales import client, ADMIN, OUTSIDER, auth
from test_operations import actor, opened, product, payload
from test_returns import setup_sale, refund
from backend.app.main import engine, User, Stock, Sale, SaleReturn, password_hash

def key(headers=ADMIN):return {**headers,'Idempotency-Key':uuid.uuid4().hex}

def test_minimum_per_branch_alerts_and_permissions():
    pid=product(stock=4);cashier,_=actor();url=f'/api/stock/{pid}/minimum'
    assert client.put(url,headers=cashier,json={'branch_id':6,'minimum':4}).status_code==403
    assert client.put(url,headers=OUTSIDER,json={'branch_id':6,'minimum':4}).status_code==404
    assert client.put(url,headers=ADMIN,json={'branch_id':6,'minimum':4}).status_code==200
    alerts=client.get('/api/stock/alerts?branch_id=6',headers=ADMIN).json()
    assert any(p['id']==pid and p['low_stock'] and p['minimum']==4 for p in alerts)
    assert not any(p['id']==pid for p in client.get('/api/stock/alerts?branch_id=5',headers=ADMIN).json())
    assert client.put(url,headers=ADMIN,json={'branch_id':6,'minimum':-1}).status_code==422
    assert client.put(url,headers=ADMIN,json={'branch_id':6,'minimum':0}).status_code==200
    assert not any(p['id']==pid for p in client.get('/api/stock/alerts?branch_id=6',headers=ADMIN).json())

def test_user_edit_password_disable_scope_and_token_revocation():
    headers,uid=actor();username='manage-'+uuid.uuid4().hex[:10]
    assert client.get('/api/users',headers=headers).status_code==403
    rows=client.get('/api/users',headers=ADMIN).json();target=next(u for u in rows if u['id']==uid)
    assert 'password_hash' not in target
    data={'username':username,'role':'almacenista','active':True,'branch_ids':[6],'password':'new-secure-password-123'}
    assert client.put(f'/api/users/{uid}',headers=OUTSIDER,json=data).status_code==404
    assert client.put(f'/api/users/{uid}',headers=ADMIN,json={**data,'branch_ids':[7]}).status_code==422
    assert client.put(f'/api/users/{uid}',headers=ADMIN,json=data).status_code==200
    assert client.get('/api/auth/me',headers=headers).status_code==401
    new=auth(username,'new-secure-password-123');assert client.get('/api/auth/me',headers=new).json()['role']=='almacenista'
    assert client.put(f'/api/users/{uid}',headers=ADMIN,json={**data,'active':False}).status_code==200
    assert client.get('/api/auth/me',headers=new).status_code==401
    me=client.get('/api/auth/me',headers=ADMIN).json()
    assert client.put(f"/api/users/{me['id']}",headers=ADMIN,json={'username':me['username'],'role':'cajero','active':True,'branch_ids':[6]}).status_code==409

def test_open_shift_blocks_user_reassignment():
    headers,uid=actor();opened(headers)
    username=client.get('/api/auth/me',headers=headers).json()['username']
    data={'username':username,'role':'cajero','active':False,'branch_ids':[6]}
    assert client.put(f'/api/users/{uid}',headers=ADMIN,json=data).status_code==409


def test_report_period_cashier_returns_profit_and_excel():
    cashier,turn,pid,sale=setup_sale();uid=client.get('/api/auth/me',headers=cashier).json()['id']
    r=client.post(f"/api/sales/{sale['id']}/returns",headers=key(),json=refund(sale,turn));assert r.status_code==201
    with Session(engine) as db:
        db.get(Sale,sale['id']).created_at=datetime(2025,1,1,18,tzinfo=timezone.utc)
        db.get(SaleReturn,r.json()['id']).created_at=datetime(2025,1,2,18,tzinfo=timezone.utc)
        db.commit()
    q=f'branch_id=6&start=2025-01-01&end=2025-01-02&cashier_id={uid}'
    response=client.get('/api/reports/sales?'+q,headers=ADMIN);assert response.status_code==200,response.text
    report=response.json();assert Decimal(report['gross'])==Decimal('.15') and Decimal(report['refunds'])==Decimal('.05')
    assert Decimal(report['total'])==Decimal('.10') and Decimal(report['cost'])==Decimal('4')
    assert Decimal(report['profit'])==Decimal('-3.90') and report['products'][0]['units']==2
    only_return=client.get(f'/api/reports/sales?branch_id=6&start=2025-01-02&end=2025-01-02&cashier_id={uid}',headers=ADMIN).json()
    assert only_return['sales_count']==0 and Decimal(only_return['total'])==Decimal('-.05')
    assert client.get('/api/reports/sales?'+q,headers=cashier).status_code==403
    assert client.get('/api/reports/sales?'+q,headers=OUTSIDER).status_code==404
    assert client.get('/api/reports/sales?branch_id=6&start=2025-01-03&end=2025-01-01',headers=ADMIN).status_code==422
    excel=client.get('/api/reports/sales.xlsx?'+q,headers=ADMIN);assert excel.status_code==200,excel.text[:100]
    wb=load_workbook(BytesIO(excel.content));assert wb.sheetnames==['Resumen','Movimientos','Productos']
    assert wb['Movimientos'].max_row==3


def test_report_timezone_boundary_and_formula_injection():
    cashier,_=actor();turn=opened(cashier);pid=product(name='=HYPERLINK("evil")',price='1',tax_exempt=True)
    sale=client.post('/api/sales',headers=key(cashier),json=payload(pid,turn)).json();uid=client.get('/api/auth/me',headers=cashier).json()['id']
    with Session(engine) as db:db.get(Sale,sale['id']).created_at=datetime(2025,1,2,5,59,tzinfo=timezone.utc);db.commit()
    q=f'branch_id=6&start=2025-01-01&end=2025-01-01&cashier_id={uid}'
    assert client.get('/api/reports/sales?'+q,headers=ADMIN).json()['sales_count']==1
    excel=client.get('/api/reports/sales.xlsx?'+q,headers=ADMIN)
    wb=load_workbook(BytesIO(excel.content));assert wb['Productos']['B2'].data_type=='s'

def test_user_duplicate_and_creation_validation():
    headers,uid=actor();username=client.get('/api/auth/me',headers=headers).json()['username']
    data={'username':'owner','role':'cajero','active':True,'branch_ids':[6]}
    assert client.put(f'/api/users/{uid}',headers=ADMIN,json=data).status_code==409
    assert client.get('/api/auth/me',headers=headers).json()['username']==username
    assert client.post('/api/users',headers=ADMIN,json={'username':'con espacios','password':'secure-password-123','role':'cajero','branch_ids':[6]}).status_code==422
